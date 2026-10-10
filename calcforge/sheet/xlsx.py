"""Reading an Excel workbook (.xlsx) into sheet sections (spreadsheet phase 6).

The user's choices (docs/SPREADSHEET_DESIGN.md): an .xlsx opens as sheet
pages, cells only — values, formulas, formatting, column widths and row
heights, merges, names, comments, data validation, conditional formatting
and each worksheet's page layout. Charts and pictures are left out (and said
so); nothing is written back to Excel.

:func:`read_workbook` gives each worksheet as the record of a sheet (the
dictionary :func:`calcforge.sheet.store.load_sheet` reads), with its paper,
and a list of what could not come across.
"""
from __future__ import annotations

import colorsys
import datetime as _dt
import re
from dataclasses import dataclass, field
from typing import Optional

from . import formula as F
from .refs import col_letters, quote_sheet

# Excel's default theme (Office 2013 and later), in the order of the theme
# index a colour names: lt1, dk1, lt2, dk2, accent1..6, hlink, folHlink
_DEFAULT_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5",
                  "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]

_PAPER = {1: "Letter", 3: "Tabloid", 5: "Legal", 8: "A3", 9: "A4", 11: "A5", 66: "A2", 67: "A3"}

_BORDERS = {"thin": "thin", "medium": "medium", "thick": "thick", "dashed": "dashed",
            "dotted": "dotted", "double": "double", "hair": "hair", "mediumDashed": "dashed",
            "dashDot": "dashed", "mediumDashDot": "dashed", "dashDotDot": "dotted",
            "mediumDashDotDot": "dotted", "slantDashDot": "dashed"}

_H_ALIGN = {"general": "general", "left": "left", "center": "center", "right": "right",
            "fill": "fill", "justify": "justify", "centerContinuous": "centerAcross",
            "distributed": "justify"}
_V_ALIGN = {"top": "top", "center": "center", "bottom": "bottom", "justify": "center",
            "distributed": "center"}

_OPS = {"between": "between", "notBetween": "notBetween", "equal": "equal",
        "notEqual": "notEqual", "greaterThan": "greater", "lessThan": "less",
        "greaterThanOrEqual": "greaterOrEqual", "lessThanOrEqual": "lessOrEqual"}


@dataclass
class ImportedSheet:
    name: str
    data: dict                         # a sheet's record (store.load_sheet)
    paper: Optional[str] = None        # "A4", "Letter"...
    orientation: Optional[str] = None  # "portrait" / "landscape"
    margins_mm: Optional[tuple] = None  # left, top, right, bottom
    hidden: bool = False
    excel_values: dict = field(default_factory=dict)   # (row, col) -> Excel's last value of a formula


@dataclass
class Imported:
    sheets: list
    left_out: list = field(default_factory=list)       # what could not come across, in words


# -- colours --------------------------------------------------------------------------------------
def _theme_colours(wb) -> list:
    xml = getattr(wb, "loaded_theme", None)
    if not xml:
        return list(_DEFAULT_THEME)
    if isinstance(xml, bytes):
        xml = xml.decode("utf-8", "replace")
    found = {}
    for key in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
                "accent5", "accent6", "hlink", "folHlink"):
        m = re.search(rf"<a:{key}>(.*?)</a:{key}>", xml, re.S)
        if not m:
            continue
        rgb = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', m.group(1))
        if rgb:
            found[key] = rgb.group(1).upper()
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    return [found.get(k, d) for k, d in zip(order, _DEFAULT_THEME)]


def _tinted(hex6: str, tint: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "".join(f"{round(v * 255):02X}" for v in (r, g, b))


def _colour(c, theme: list) -> Optional[str]:
    """An openpyxl colour as #RRGGBB (None: automatic or none)."""
    if c is None:
        return None
    try:
        kind = c.type
    except AttributeError:
        return None
    hex6 = None
    if kind == "rgb" and isinstance(c.rgb, str):
        argb = c.rgb
        if len(argb) == 8:
            hex6 = argb[2:]
        elif len(argb) == 6:
            hex6 = argb
    elif kind == "theme" and c.theme is not None and 0 <= c.theme < len(theme):
        hex6 = theme[c.theme]
    elif kind == "indexed" and c.indexed is not None:
        from openpyxl.styles.colors import COLOR_INDEX
        if c.indexed in (64, 65):
            return None                      # system foreground / background: automatic
        if 0 <= c.indexed < len(COLOR_INDEX):
            hex6 = COLOR_INDEX[c.indexed][2:]
    if hex6 is None:
        return None
    tint = getattr(c, "tint", 0.0) or 0.0
    if tint:
        hex6 = _tinted(hex6, tint)
    return "#" + hex6.upper()


# -- cells -----------------------------------------------------------------------------------------
def _style(cell, theme, default_font) -> dict:
    out = {}
    fmt = cell.number_format
    if fmt and fmt != "General":
        out["number_format"] = fmt
    font = cell.font
    if font is not None:
        if font.name and font.name != default_font[0]:
            out["font"] = font.name
        if font.sz and float(font.sz) != default_font[1]:
            out["size"] = float(font.sz)
        if font.b:
            out["bold"] = True
        if font.i:
            out["italic"] = True
        if font.u:
            out["underline"] = "double" if "double" in str(font.u) else "single"
        if font.strike:
            out["strike"] = True
        colour = _colour(font.color, theme)
        if colour and colour != "#000000":
            out["color"] = colour
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) not in (None, "none"):
        colour = _colour(fill.fgColor, theme) if fill.fill_type == "solid" else \
            (_colour(fill.fgColor, theme) or _colour(fill.bgColor, theme))
        if colour:
            out["fill"] = colour
    border = cell.border
    if border is not None:
        for side in ("left", "right", "top", "bottom"):
            s = getattr(border, side, None)
            if s is not None and s.style in _BORDERS:
                out[side] = [_BORDERS[s.style], _colour(s.color, theme) or "#000000"]
    al = cell.alignment
    if al is not None:
        if al.horizontal and _H_ALIGN.get(al.horizontal, "general") != "general":
            out["h_align"] = _H_ALIGN[al.horizontal]
        if al.vertical and _V_ALIGN.get(al.vertical, "bottom") != "bottom":
            out["v_align"] = _V_ALIGN[al.vertical]
        if al.wrap_text:
            out["wrap"] = True
        if al.shrink_to_fit:
            out["shrink"] = True
        if al.indent:
            out["indent"] = int(al.indent)
        if al.text_rotation:
            rot = int(al.text_rotation)
            out["rotation"] = 255 if rot == 255 else (rot if rot <= 90 else 90 - rot)
    pro = cell.protection
    if pro is not None:
        if pro.locked is False:
            out["locked"] = False
        if pro.hidden:
            out["hidden"] = True
    return out


_EPOCH = _dt.datetime(1899, 12, 30)


def _serial(v) -> float:
    if isinstance(v, _dt.datetime):
        return (v - _EPOCH).total_seconds() / 86400.0
    if isinstance(v, _dt.date):
        return float((v - _EPOCH.date()).days)
    if isinstance(v, _dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6) / 86400.0
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 86400.0
    return float(v)


def _number_text(x: float) -> str:
    from .values import general_number
    if x == int(x) and abs(x) < 1e15:
        return str(int(x))
    text = repr(float(x))
    return text if "e" not in text else general_number(x)


def _value_input(v, data_type: str) -> Optional[str]:
    """What to type to get an Excel value back."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return _number_text(float(v))
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return _number_text(_serial(v))
    text = str(v)
    if data_type == "e":
        return text
    if not text:
        return None
    from .inputs import read_value
    got, _fmt = read_value(text)
    # text that would be read as something else stays text
    if not (isinstance(got, str) and got == text) or text.startswith(("=", "'")):
        return "'" + text
    return text


# -- formulas --------------------------------------------------------------------------------------
_ANCHOR = re.compile(r"_xlfn\.ANCHORARRAY\(\s*((?:'[^']+'|[A-Za-z0-9_.]+)?!?\$?[A-Za-z]{1,3}\$?\d+)\s*\)",
                     re.I)


def _plain_formula(text: str) -> str:
    """Excel's stored formula as it is typed: no _xlfn./_xlpm. prefixes,
    ANCHORARRAY(A1) as A1#, SINGLE(x) (the @ operator) as x."""
    text = _ANCHOR.sub(lambda m: m.group(1) + "#", text)
    text = re.sub(r"_xlfn\.SINGLE\(", "@(", text, flags=re.I)
    text = re.sub(r"_xlpm\.", "", text, flags=re.I)
    text = re.sub(r"_xlfn\._xlws\.", "", text, flags=re.I)
    text = re.sub(r"_xlfn\.", "", text, flags=re.I)
    text = re.sub(r"_xlws\.", "", text, flags=re.I)
    return text


@dataclass
class _Table:
    """An Excel table (ListObject) on a worksheet."""

    name: str
    sheet: str
    top: int
    left: int
    bottom: int
    right: int
    header_rows: int
    totals_rows: int
    columns: list


def _tables_of(wb) -> dict:
    out = {}
    for ws in wb.worksheets:
        for table in getattr(ws, "tables", {}).values():
            name = table.displayName or table.name
            from openpyxl.utils import range_boundaries
            left, top, right, bottom = range_boundaries(table.ref)
            cols = [c.name for c in table.tableColumns] if table.tableColumns else []
            out[name.lower()] = _Table(name, ws.title, top - 1, left - 1, bottom - 1, right - 1,
                                       int(table.headerRowCount if table.headerRowCount is not None else 1),
                                       int(table.totalsRowCount or 0), cols)
    return out


def _structured_to_a1(text: str, sheet_name: str, row: int, col: int, tables: dict) -> str:
    """Excel's table references (Loads[Load], [@Span]...) written as the
    cells they name: a worksheet's Excel tables are not tables here."""
    if "[" not in text or not tables:
        return text
    try:
        tokens = F.tokenize(text)
    except F.FormulaError:
        return text
    changed = False
    for t in tokens:
        if t.kind != "struct":
            continue
        spec = t.ref
        table = tables.get(spec.table.lower()) if spec.table else next(
            (tb for tb in tables.values() if tb.sheet == sheet_name
             and tb.top <= row <= tb.bottom and tb.left <= col <= tb.right), None)
        if table is None:
            continue
        lower = [c.lower() for c in table.columns]
        if spec.first is None:
            left, right = table.left, table.right
        else:
            a = lower.index(spec.first.lower()) if spec.first.lower() in lower else None
            b = lower.index(spec.last.lower()) if spec.last and spec.last.lower() in lower else a
            if a is None or b is None:
                continue
            left, right = table.left + min(a, b), table.left + max(a, b)
        data_top = table.top + table.header_rows
        data_bottom = table.bottom - table.totals_rows
        specials = set(spec.specials) or {"#Data"}
        if "@" in specials:
            top = bottom = row
            ref = (f"${col_letters(left)}{top + 1}" if left == right else
                   f"${col_letters(left)}{top + 1}:${col_letters(right)}{bottom + 1}")
        else:
            if "#All" in specials:
                top, bottom = table.top, table.bottom
            elif "#Headers" in specials and "#Data" in specials:
                top, bottom = table.top, data_bottom
            elif "#Data" in specials and "#Totals" in specials:
                top, bottom = data_top, table.bottom
            elif "#Headers" in specials:
                top = bottom = table.top
            elif "#Totals" in specials:
                top = bottom = table.bottom
            else:
                top, bottom = data_top, data_bottom
            ref = f"${col_letters(left)}${top + 1}"
            if (top, left) != (bottom, right):
                ref += f":${col_letters(right)}${bottom + 1}"
        if table.sheet != sheet_name:
            ref = quote_sheet(table.sheet) + "!" + ref
        t.text = ref
        changed = True
    return F.join(tokens) if changed else text


# -- the workbook -----------------------------------------------------------------------------------
def _width_pt(chars: float) -> float:
    """Excel's column width (characters of Calibri 11) in points."""
    px = int((256 * chars + int(128 / 7)) / 256 * 7)
    return px * 0.75


def read_workbook(path: str, taken: Optional[set] = None) -> Imported:
    """Every worksheet of an .xlsx as a sheet's record; *taken*: the names
    already used in the document (a worksheet so named is renamed, and the
    formulas reading it follow)."""
    import openpyxl
    from openpyxl.utils import range_boundaries

    wb = openpyxl.load_workbook(path, data_only=False, rich_text=False)
    try:
        cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        cached = None
    theme = _theme_colours(wb)
    try:
        base = wb._fonts[0]
        default_font = (base.name or "Calibri", float(base.sz or 11))
    except (AttributeError, IndexError, TypeError):
        default_font = ("Calibri", 11.0)
    left_out: list = []
    taken = {t.lower() for t in (taken or set())}
    renames = {}
    for ws in wb.worksheets:
        name = ws.title
        if name.lower() in taken:
            k = 2
            while f"{name} ({k})".lower() in taken:
                k += 1
            renames[name] = f"{name} ({k})"
            taken.add(renames[name].lower())
        else:
            taken.add(name.lower())
    for cs in getattr(wb, "chartsheets", []):
        left_out.append(f"chart sheet “{cs.title}”")
    tables = _tables_of(wb)

    def renamed(text: str) -> str:
        for old, new in renames.items():
            text = F.rename_sheet_in(text, old, new)
        return text

    sheets = []
    for ws in wb.worksheets:
        name = renames.get(ws.title, ws.title)
        styles = [{}]
        style_index: dict = {}
        cells = []
        excel_values = {}
        cached_ws = cached[ws.title] if cached is not None and ws.title in cached.sheetnames else None
        cached_cells = {}
        if cached_ws is not None:
            for row in cached_ws.iter_rows():
                for c in row:
                    if c.value is not None and hasattr(c, "row"):
                        cached_cells[(c.row - 1, c.column - 1)] = c.value
        spilled = set()
        for row in ws.iter_rows():
            for c in row:
                r, k = c.row - 1, c.column - 1
                value = c.value
                text = None
                if value is not None and type(value).__name__ == "ArrayFormula":
                    # a dynamic array (or an old Ctrl+Shift+Enter one): its
                    # formula in its first cell spills over the rest
                    text = "=" + _plain_formula(str(value.text).lstrip("="))
                    try:
                        l2, t2, r2, b2 = range_boundaries(value.ref)
                        for rr in range(t2 - 1, b2):
                            for cc in range(l2 - 1, r2):
                                if (rr, cc) != (r, k):
                                    spilled.add((rr, cc))
                    except (ValueError, TypeError):
                        pass
                elif value is not None and type(value).__name__ == "DataTableFormula":
                    left_out.append(f"a What-If data table on “{ws.title}” (its values kept)")
                    text = _value_input(cached_cells.get((r, k)), "n")
                elif c.data_type == "f" or (isinstance(value, str) and value.startswith("=")
                                            and c.data_type != "s"):
                    text = "=" + _plain_formula(str(value)[1:])
                else:
                    text = _value_input(value, c.data_type)
                if (r, k) in spilled:
                    text = None
                if text is not None and text.startswith("="):
                    text = "=" + _structured_to_a1(text[1:], ws.title, r, k, tables)
                    text = "=" + renamed(text[1:])
                    if (r, k) in cached_cells:
                        excel_values[(r, k)] = cached_cells[(r, k)]
                st = _style(c, theme, default_font) if c.has_style else {}
                comment = c.comment.text if c.comment is not None else None
                if text is None and not st and not comment:
                    continue
                key = tuple(sorted((a, tuple(b) if isinstance(b, list) else b) for a, b in st.items()))
                index = style_index.get(key)
                if index is None:
                    index = len(styles) if st else 0
                    if st:
                        styles.append(st)
                    style_index[key] = index
                entry = [r, k, text or "", index]
                if comment:
                    entry.append(comment.strip())
                cells.append(entry)
        # column widths and row heights
        fmt = ws.sheet_format
        default_width = _width_pt(float(fmt.defaultColWidth)) if fmt.defaultColWidth else \
            (_width_pt(float(fmt.baseColWidth) + 0.71) if fmt.baseColWidth not in (None, 8) else 48.0)
        default_height = float(fmt.defaultRowHeight or 15.0)
        widths, hidden_cols = {}, []
        for key, dim in ws.column_dimensions.items():
            lo, hi = (dim.min or 0), (dim.max or 0)
            if not lo:
                from openpyxl.utils import column_index_from_string
                lo = hi = column_index_from_string(key)
            for k in range(lo - 1, min(hi, lo + 1000)):
                if dim.width and dim.customWidth is not False:
                    widths[str(k)] = round(_width_pt(float(dim.width)), 2)
                if dim.hidden:
                    hidden_cols.append(k)
        heights, hidden_rows = {}, []
        for r, dim in ws.row_dimensions.items():
            if dim.ht is not None and dim.customHeight:
                heights[str(r - 1)] = float(dim.ht)
            if dim.hidden:
                hidden_rows.append(r - 1)
        merges = []
        for mr in ws.merged_cells.ranges:
            merges.append([mr.min_row - 1, mr.min_col - 1, mr.max_row - 1, mr.max_col - 1])
        page = _page_options(ws, renamed)
        data = {
            "name": name, "kind": "sheet", "size": None, "cells": cells, "styles": styles,
            "widths": widths, "heights": heights, "hidden_rows": sorted(set(hidden_rows)),
            "hidden_cols": sorted(set(hidden_cols)), "merges": merges,
            "default_width": round(default_width, 2), "default_height": default_height,
            "gridlines": bool(ws.sheet_view.showGridLines if ws.sheet_view.showGridLines is not None
                              else True),
            "headings": bool(ws.print_options.headings) if ws.print_options else False,
            "names": [], "cond_rules": _cond_rules(ws, theme, renamed),
            "validations": _validations(ws, renamed), "filter": None, "filtered_rows": [],
            "page": page,
        }
        setup = ws.page_setup
        paper = _PAPER.get(int(setup.paperSize)) if setup is not None and setup.paperSize else None
        orientation = setup.orientation if setup is not None and setup.orientation in (
            "portrait", "landscape") else None
        m = ws.page_margins
        margins = (m.left * 25.4, m.top * 25.4, m.right * 25.4, m.bottom * 25.4) if m is not None else None
        if getattr(ws, "_charts", None):
            left_out.append(f"{len(ws._charts)} chart{'s' * (len(ws._charts) > 1)} on “{ws.title}”")
        if getattr(ws, "_images", None):
            left_out.append(f"{len(ws._images)} picture{'s' * (len(ws._images) > 1)} on “{ws.title}”")
        if setup is not None and setup.scale and int(setup.scale) != 100 and not page.get("fit_width"):
            left_out.append(f"the print scale of {int(setup.scale)}% on “{ws.title}”")
        sheets.append(ImportedSheet(name, data, paper, orientation, margins,
                                    ws.sheet_state != "visible", excel_values))
    _names(wb, sheets, renames, renamed, left_out)
    return Imported(sheets, left_out)


def _page_options(ws, renamed) -> dict:
    from openpyxl.utils import range_boundaries
    out = {}
    try:
        area = ws.print_area
    except Exception:  # noqa: BLE001
        area = None
    if area:
        first = str(area).split(",")[0]
        ref = first.split("!")[-1].replace("$", "")
        try:
            l, t, r, b = range_boundaries(ref)
            out["print_area"] = [t - 1, l - 1, b - 1, r - 1]
        except (ValueError, TypeError):
            pass
    titles = ws.print_title_rows
    if titles:
        m = re.match(r"\$?(\d+):\$?(\d+)", str(titles).split("!")[-1])
        if m:
            out["titles"] = [int(m.group(1)) - 1, int(m.group(2)) - 1]
    breaks = [int(b.id) for b in (ws.row_breaks.brk if ws.row_breaks else []) if b.id]
    if breaks:
        out["breaks"] = sorted(set(breaks))
    pr = ws.sheet_properties.pageSetUpPr if ws.sheet_properties is not None else None
    if pr is not None and pr.fitToPage and ws.page_setup.fitToWidth in (1, None):
        out["fit_width"] = True
    po = ws.print_options
    if po is not None:
        if po.horizontalCentered:
            out["center_h"] = True
        if po.verticalCentered:
            out["center_v"] = True
        if po.gridLines:
            out["print_gridlines"] = True
        if po.headings:
            out["print_headings"] = True
    return out


def _cond_rules(ws, theme, renamed) -> list:
    from openpyxl.utils import range_boundaries
    out = []
    try:
        blocks = list(ws.conditional_formatting)
    except Exception:  # noqa: BLE001
        return out
    for cf in blocks:
        ranges = []
        for part in str(cf.sqref).split():
            try:
                l, t, r, b = range_boundaries(part)
                ranges.append([t - 1, l - 1, b - 1, r - 1])
            except (ValueError, TypeError):
                continue
        if not ranges:
            continue
        for rule in cf.rules:
            got = _rule(rule, theme, renamed)
            if got is not None:
                got["ranges"] = [list(x) for x in ranges]
                got["stop"] = bool(rule.stopIfTrue)
                out.append((rule.priority or 0, len(out), got))
    # Excel's priority runs across the whole sheet, not within each range
    return [got for _p, _i, got in sorted(out, key=lambda x: x[:2])]


def _dxf_format(dxf, theme) -> dict:
    fmt = {}
    if dxf is None:
        return fmt
    if dxf.font is not None:
        c = _colour(dxf.font.color, theme)
        if c:
            fmt["color"] = c
        if dxf.font.b:
            fmt["bold"] = True
        if dxf.font.i:
            fmt["italic"] = True
        if dxf.font.u:
            fmt["underline"] = True
        if dxf.font.strike:
            fmt["strike"] = True
    if dxf.fill is not None:
        c = _colour(getattr(dxf.fill, "bgColor", None), theme) or \
            _colour(getattr(dxf.fill, "fgColor", None), theme)
        if c:
            fmt["fill"] = c
    if dxf.numFmt is not None and dxf.numFmt.formatCode:
        fmt["number_format"] = dxf.numFmt.formatCode
    return fmt


def _operand(text) -> str:
    text = str(text)
    try:
        float(text)
        return text
    except ValueError:
        pass
    if text.startswith('"') and text.endswith('"'):
        return text[1:-1]
    return "=" + text


def _rule(rule, theme, renamed) -> Optional[dict]:
    kind = rule.type
    fmt = _dxf_format(rule.dxf, theme)
    formulas = [renamed(_plain_formula(str(f))) for f in (rule.formula or [])]
    if kind == "cellIs" and rule.operator in _OPS and formulas:
        out = {"type": "cell", "op": _OPS[rule.operator], "a": _operand(formulas[0]), "format": fmt}
        if len(formulas) > 1:
            out["b"] = _operand(formulas[1])
        return out
    if kind == "expression" and formulas:
        return {"type": "formula", "formula": "=" + formulas[0], "format": fmt}
    if kind in ("containsText", "notContainsText", "beginsWith", "endsWith"):
        how = {"containsText": "contains", "notContainsText": "notContains",
               "beginsWith": "begins", "endsWith": "ends"}[kind]
        return {"type": "text", "how": how, "text": rule.text or "", "format": fmt}
    simple = {"containsBlanks": "blanks", "notContainsBlanks": "noBlanks",
              "containsErrors": "errors", "notContainsErrors": "noErrors",
              "duplicateValues": "duplicate", "uniqueValues": "unique"}
    if kind in simple:
        return {"type": simple[kind], "format": fmt}
    if kind == "top10":
        return {"type": "bottom" if rule.bottom else "top", "n": int(rule.rank or 10),
                "percent": bool(rule.percent), "format": fmt}
    if kind == "aboveAverage":
        above = rule.aboveAverage is not False
        return {"type": "above" if above else "below", "equal": bool(rule.equalAverage),
                "format": fmt}
    if kind == "timePeriod" and rule.timePeriod:
        when = {"today": "today", "yesterday": "yesterday", "tomorrow": "tomorrow",
                "last7Days": "last7", "thisMonth": "thisMonth", "lastMonth": "lastMonth",
                "nextMonth": "nextMonth"}.get(rule.timePeriod)
        if when:
            return {"type": "date", "when": when, "format": fmt}
    if kind == "colorScale" and rule.colorScale is not None:
        colours = [_colour(c, theme) or "#FFFFFF" for c in rule.colorScale.color]
        return {"type": "scale", "colors": colours}
    if kind == "dataBar" and rule.dataBar is not None:
        return {"type": "databar", "color": _colour(rule.dataBar.color, theme) or "#638EC6"}
    if kind == "iconSet" and rule.iconSet is not None:
        name = str(rule.iconSet.iconSet or "3TrafficLights1")
        family = "arrows" if "Arrow" in name else "flags" if "Flag" in name else \
            "symbols" if "Symbol" in name else "traffic"
        return {"type": "icons", "set": family, "reverse": bool(rule.iconSet.reverse)}
    return None


def _validations(ws, renamed) -> list:
    from openpyxl.utils import range_boundaries
    out = []
    dv_list = getattr(ws, "data_validations", None)
    for dv in (dv_list.dataValidation if dv_list is not None else []):
        ranges = []
        for part in str(dv.sqref).split():
            try:
                l, t, r, b = range_boundaries(part)
                ranges.append([t - 1, l - 1, b - 1, r - 1])
            except (ValueError, TypeError):
                continue
        kind = {"whole": "whole", "decimal": "decimal", "list": "list", "date": "date",
                "time": "time", "textLength": "length", "custom": "custom"}.get(dv.type or "", None)
        if not ranges or kind is None:
            continue
        rule = {"ranges": ranges, "type": kind, "blank": bool(dv.allow_blank),
                # Excel's showDropDown means the opposite of what it says
                "dropdown": not bool(dv.showDropDown),
                "input_title": dv.promptTitle or "", "input": dv.prompt or "",
                "style": {"stop": "stop", "warning": "warning", "information": "information"}.get(
                    dv.errorStyle or "stop", "stop"),
                "error_title": dv.errorTitle or "", "error": dv.error or ""}
        f1 = renamed(_plain_formula(dv.formula1)) if dv.formula1 else ""
        f2 = renamed(_plain_formula(dv.formula2)) if dv.formula2 else ""
        if kind == "list":
            rule["source"] = f1[1:-1] if f1.startswith('"') else "=" + f1
        elif kind == "custom":
            rule["formula"] = "=" + f1
        else:
            rule["op"] = _OPS.get(dv.operator or "between", "between")
            rule["a"] = _operand(f1) if f1 else ""
            if f2:
                rule["b"] = _operand(f2)
        out.append(rule)
    return out


def _names(wb, sheets: list, renames: dict, renamed, left_out: list) -> None:
    """The workbook's defined names, each kept with the sheet whose cells it
    names (or the first sheet, for a constant)."""
    by_name = {s.name.lower(): s for s in sheets}
    entries = []
    try:
        for name, dn in wb.defined_names.items():
            entries.append((name, dn, None))
    except AttributeError:
        pass
    for ws in wb.worksheets:
        for name, dn in getattr(ws, "defined_names", {}).items():
            entries.append((name, dn, renames.get(ws.title, ws.title)))
    for name, dn, local in entries:
        if name.startswith("_xlnm.") or getattr(dn, "hidden", False):
            continue
        text = str(dn.attr_text or "")
        if "#REF!" in text:
            left_out.append(f"the name “{name}” (it refers to cells that are gone)")
            continue
        text = renamed(_plain_formula(text))
        try:
            tokens = F.tokenize(text)
        except F.FormulaError:
            left_out.append(f"the name “{name}”")
            continue
        first = next((t for t in tokens if t.kind == "ref" and t.ref.sheet), None)
        home = by_name.get(first.ref.sheet.lower()) if first is not None else None
        if home is None:
            home = by_name.get(local.lower()) if local else (sheets[0] if sheets else None)
        if home is None:
            continue
        home.data["names"].append([name, text, local is not None and local.lower() == home.name.lower(), ""])
