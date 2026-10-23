"""Random sessions in the real window, looking for what breaks (the user,
2026-10-24: "I asked you to find the bugs, not me … all features of app
test in all possible ways").

A seeded run does what a person does, in any order, for hundreds of steps:
clicks cells and types numbers, text, units and formulas (across sheets,
by name, lookups, errors); inserts and deletes rows and columns; drags cells
by their border and the fill handle; copies, cuts and pastes; renames and
adds and deletes sheets; freezes panes; drags the print area and the page
breaks; zooms and scrolls; goes in and out of the workbook view; adds,
duplicates, moves and deletes pages; brings in an Excel workbook (charts,
pictures, names Excel left as #REF!); draws markups on pages and sheets,
moves and deletes them; types equations that read the sheets; runs menu
commands; undoes and redoes.

After every step it checks:

* nothing raised (an exception anywhere, in any event handler);
* the document holds together: a frame for every page, a thumbnail for every
  page, every sheet's cells on its first page, every formula and name
  readable;
* now and then, that undoing the step puts the document back exactly as it
  was, and redoing it gives the step's result again;
* now and then, that saving and opening the document again gives the same
  document.

A failure is written with its seed and step, so it replays:

    python tools/qa/monkey.py run --seeds 1-8 --steps 400 [--out DIR]
    python tools/qa/monkey.py replay --seed 5 --upto 212 [--verbose]
"""
from __future__ import annotations

import argparse
import json
import os
import random
import sys
import tempfile
import time
import traceback

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)

from tools.qa import exercise as X  # noqa: E402

NAMES = ["Loads", "Data", "Calc", "Sheet 2", "Loads (2)", "Tables", "Inputs", "x1"]
TEXTS = ["Load", "Span", "Beam B1", "N/A", "kN", "TRUE", "  ", "Total"]
UNITS = ["5 kN", "2.5 m", "12 kN*m", "3'm", "0.5 MPa", "100 mm"]


class Session:
    def __init__(self, seed: int, out: str, verbose: bool = False):
        self.seed = seed
        self.rng = random.Random(seed)
        self.out = out
        self.verbose = verbose
        self.collect = bool(os.environ.get("MONKEY_GC"))
        self.log: list = []
        self.failures: list = []
        self.workbooks: list = []
        X.app()
        self.win = X.new_window()
        self.win.interactive_prompts = False
        X.Watch.policy = "cancel"
        self.tmp = tempfile.mkdtemp(prefix=f"monkey-{seed}-")
        self._make_workbooks()

    # -- what is at hand ---------------------------------------------------------------
    @property
    def view(self):
        return self.win.view

    def runs(self) -> list:
        from calcforge.items.sheetpage import SheetRunItem
        out = []
        for frame in getattr(self.win.scene, "frames", []) or []:
            for item in frame.markups():
                if isinstance(item, SheetRunItem) and item.sheet is not None:
                    out.append(item)
        return out

    def shown(self):
        from calcforge.ui import sheetview
        uid = sheetview.run_shown(self.win.scene)
        return next((r for r in self.runs() if r.uid == uid), None) if uid else None

    def sheet_names(self) -> list:
        return [r.sheet.name for r in self.runs()]

    def pick(self, seq):
        return seq[self.rng.randrange(len(seq))] if seq else None

    def _make_workbooks(self) -> None:
        try:
            import openpyxl
            from openpyxl.chart import BarChart, Reference, ScatterChart, Series
            from openpyxl.workbook.defined_name import DefinedName
        except ImportError:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calculations"
        for r in range(1, 12):
            ws.cell(r, 1, r)
            ws.cell(r, 2, f"=A{r}*2")
            ws.cell(r, 3, f"=B{r}+Tables!A1")
        ws["E1"] = "=INDEX(A1:C1,,2)"
        ws["E2"] = "=SUM(A1:A11)"
        ws["E3"] = "=VLOOKUP(3,A1:B11,2,FALSE)"
        ws["E4"] = "=Z*2"
        tables = wb.create_sheet("Tables")
        tables["A1"] = 10
        tables["B1"] = "=Calculations!A3"
        wb.defined_names["Z"] = DefinedName("Z", attr_text="Tables!$A$1")
        wb.defined_names["Gone"] = DefinedName("Gone", attr_text="Tables!#REF!")
        chart = ScatterChart()
        chart.title = "Coefficient"
        chart.series.append(Series(Reference(ws, min_col=2, min_row=1, max_row=11),
                                   Reference(ws, min_col=1, min_row=1, max_row=11), title="B"))
        ws.add_chart(chart, "G2")
        bars = BarChart()
        bars.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
        ws.add_chart(bars, "G20")
        path = os.path.join(self.tmp, "Parts.xlsx")
        wb.save(path)
        self.workbooks.append(path)

    # -- the actions -------------------------------------------------------------------
    def actions(self) -> list:
        sheet_on = self.shown() is not None
        acts = [
            (6, "add page", self.a_add_page),
            (5, "insert spreadsheet", self.a_insert_spreadsheet),
            (3, "insert workbook", self.a_insert_workbook),
            (3, "duplicate page", self.a_duplicate_page),
            (3, "delete page (thumbnails Delete)", self.a_delete_page),
            (3, "move page", self.a_move_page),
            (5, "go to page (thumbnail click)", self.a_thumbnail_click),
            (4, "open a sheet", self.a_enter_sheet),
            (6, "draw markup", self.a_draw_markup),
            (4, "move a markup", self.a_move_markup),
            (3, "delete a markup", self.a_delete_markup),
            (4, "equation", self.a_equation),
            (6, "undo", self.a_undo),
            (4, "redo", self.a_redo),
            (3, "zoom", self.a_zoom),
            (3, "scroll", self.a_scroll),
            (3, "menu command", self.a_menu_command),
            (2, "calc mode on/off", self.a_calc_mode),
            (3, "click-placed tool", self.a_click_tool),
            (3, "text box", self.a_text_box),
            (4, "pick a markup and press keys", self.a_markup_keys),
            (3, "resize a markup by a handle", self.a_resize_markup),
            (3, "edit an equation", self.a_edit_equation),
            (4, "right-click menu entry", self.a_context_entry),
            (3, "style a markup", self.a_style),
            (1, "select all and delete", self.a_select_all_delete),
        ]
        if sheet_on or self.runs():
            acts += [
                (14, "click a cell", self.a_click_cell),
                (14, "type into a cell", self.a_type),
                (4, "clear cells (Delete)", self.a_clear),
                (4, "insert rows/cols", self.a_insert_rc),
                (4, "delete rows/cols", self.a_delete_rc),
                (4, "drag cells by the border", self.a_move_cells),
                (3, "fill handle", self.a_fill),
                (4, "copy/cut and paste", self.a_copy_paste),
                (2, "rename sheet", self.a_rename_sheet),
                (2, "new sheet (+)", self.a_new_sheet),
                (1, "delete sheet", self.a_delete_sheet),
                (2, "define name", self.a_define_name),
                (2, "freeze panes", self.a_freeze),
                (3, "drag print area / break", self.a_drag_area),
                (3, "leave the sheet", self.a_leave_sheet),
                (2, "select rows by heading", self.a_heading_select),
                (3, "format cells (keys)", self.a_format_cells),
            ]
        return acts

    def choose(self):
        acts = self.actions()
        total = sum(w for w, _n, _f in acts)
        k = self.rng.uniform(0, total)
        for weight, name, fn in acts:
            k -= weight
            if k <= 0:
                return name, fn
        return acts[-1][1], acts[-1][2]

    # pages
    def a_add_page(self):
        doc = self.win.document
        index = self.rng.randrange(len(doc.pages))
        self.win.add_page(index, before=self.rng.random() < 0.3)
        return f"after page {index}"

    def a_insert_spreadsheet(self):
        doc = self.win.document
        index = self.rng.randrange(len(doc.pages))
        before = self.rng.random() < 0.3
        self.win.insert_spreadsheet(index, before=before)
        return f"at page {index} before={before}"

    def a_insert_workbook(self):
        if not self.workbooks:
            return "no openpyxl"
        from calcforge.ui import xlsximport
        doc = self.win.document
        index = self.rng.randrange(len(doc.pages))
        xlsximport.insert_workbook(self.win, self.workbooks[0], index)
        return f"after page {index}"

    def a_duplicate_page(self):
        index = self.rng.randrange(len(self.win.document.pages))
        self.win.duplicate_page(index)
        return f"page {index}"

    def a_delete_page(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        doc = self.win.document
        if len(doc.pages) < 2:
            return "one page only"
        strip = self.win.pages_panel.list
        rows = strip.shown_rows()
        row = self.pick(rows)
        rect = strip.visualItemRect(strip.item(row))
        strip.scrollToItem(strip.item(row))
        X.pump()
        rect = strip.visualItemRect(strip.item(row))
        QTest.mouseClick(strip.viewport(), Qt.LeftButton, Qt.NoModifier, rect.center() + _pt(0, -rect.height() // 4))
        X.pump()
        strip.setFocus()
        QTest.keyClick(strip, Qt.Key_Delete)
        return f"thumbnail row {row}"

    def a_move_page(self):
        doc = self.win.document
        if len(doc.pages) < 2:
            return "one page only"
        src = self.rng.randrange(len(doc.pages))
        target = self.rng.randrange(len(doc.pages) + 1)
        self.win.move_page(src, target, 1)
        return f"{src} -> {target}"

    def a_thumbnail_click(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        strip = self.win.pages_panel.list
        rows = strip.shown_rows()
        if not rows:
            return "none"
        row = self.pick(rows)
        strip.scrollToItem(strip.item(row))
        X.pump()
        rect = strip.visualItemRect(strip.item(row))
        QTest.mouseClick(strip.viewport(), Qt.LeftButton, Qt.NoModifier, rect.center())
        return f"row {row}"

    def a_enter_sheet(self):
        from calcforge.ui import sheetview
        runs = self.runs()
        if not runs:
            return "no sheets"
        run = self.pick(runs)
        sheetview.enter(self.win, run.uid)
        return run.sheet.name

    def a_leave_sheet(self):
        from calcforge.ui import sheetview
        if self.shown() is None:
            return "not shown"
        if self.rng.random() < 0.5:
            from PySide6.QtCore import Qt
            from PySide6.QtTest import QTest
            QTest.mouseClick(self.win.sheet_tabs.close_button, Qt.LeftButton)
            return "the ×"
        sheetview.leave(self.win)
        return "leave()"

    # cells
    def _run_for_cells(self):
        run = self.shown()
        if run is None:
            runs = self.runs()
            if not runs:
                return None
            run = self.pick(runs)
            from calcforge.ui import sheetview
            sheetview.enter(self.win, run.uid)
            X.pump(3)
            run = self.shown() or run
        return run

    def _cell_point(self, run, r, c):
        from PySide6.QtCore import QPoint
        self.view.centerOn(run.mapToScene(run.cell_rect(r, c).center()))
        X.pump()
        p = self.view.mapFromScene(run.mapToScene(run.cell_rect(r, c).center()))
        return QPoint(p.x(), p.y())

    def _click(self, point, mods=None):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        QTest.mouseClick(self.view.viewport(), Qt.LeftButton, mods or Qt.NoModifier, point)

    def a_click_cell(self):
        from PySide6.QtCore import Qt
        run = self._run_for_cells()
        if run is None:
            return "none"
        r, c = self.rng.randrange(0, 40), self.rng.randrange(0, 10)
        mods = Qt.ShiftModifier if self.rng.random() < 0.2 else None
        self._click(self._cell_point(run, r, c), mods)
        return f"{run.sheet.name} ({r},{c}){' shift' if mods else ''}"

    def _formula(self) -> str:
        names = self.sheet_names()
        other = self.pick(names) or "Sheet1"
        from calcforge.sheet.refs import quote_sheet
        q = quote_sheet(other)
        r1, r2 = self.rng.randrange(1, 20), self.rng.randrange(1, 30)
        pool = [
            f"=A{r1}+B{r2}", f"=SUM(A1:A{r2})", f"={q}!A{r1}*2", f"=AVERAGE({q}!B1:B{r2})",
            f"=IF(A{r1}>3,\"big\",\"small\")", "=1/0", f"=INDEX(A1:C1,,{self.rng.randrange(1, 4)})",
            f"=VLOOKUP({r1},A1:B{r2},2,FALSE)", f"=A{r1}&\" kN\"", "=Z*2", "=Gone+1",
            f"=MATCH({r1},A1:A{r2})", f"=ROUND(A{r1}/3,2)", f"=A{r1}'m*2", "=NOW()-NOW()",
            f"=SUM({q}!A:A)", f"=COUNTIF(A1:A{r2},\">2\")", "=", "=SUM(", "=A1+",
            f"=TEXT(A{r1},\"0.00\")", f"=CONCATENATE(\"1/\",A{r1})", f"=A{r1}#",
            f"=SEQUENCE({self.rng.randrange(1, 4)})", f"=MAX(A1:C{r2})-MIN(A1:C{r2})",
        ]
        return self.pick(pool)

    def a_type(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        run = self._run_for_cells()
        if run is None:
            return "none"
        r, c = self.rng.randrange(0, 30), self.rng.randrange(0, 8)
        self._click(self._cell_point(run, r, c))
        X.pump()
        k = self.rng.random()
        if k < 0.35:
            text = str(self.pick([1, 2.5, -3, 0, 1e6, 42, 7, 3.14159]))
        elif k < 0.5:
            text = self.pick(TEXTS)
        elif k < 0.6:
            text = self.pick(UNITS)
        else:
            text = self._formula()
        target = self.view.viewport()
        QTest.keyClicks(target, text)
        X.pump()
        editor = self.view.tables.editor
        if editor is not None:
            QTest.keyClick(editor, self.pick([Qt.Key_Return, Qt.Key_Tab, Qt.Key_Return, Qt.Key_Escape]))
        return f"{run.sheet.name} ({r},{c}) {text!r}"

    def a_clear(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        run = self._run_for_cells()
        if run is None or self.view.tables.item is None:
            return "none"
        r, c = self.rng.randrange(0, 20), self.rng.randrange(0, 6)
        self.view.tables.select((r, c), (r + self.rng.randrange(0, 4), c + self.rng.randrange(0, 3)))
        QTest.keyClick(self.view.viewport(), Qt.Key_Delete)
        return f"from ({r},{c})"

    def _select_some(self):
        run = self._run_for_cells()
        tables = self.view.tables
        if run is None or tables.item is None:
            return None
        r, c = self.rng.randrange(0, 20), self.rng.randrange(0, 6)
        tables.select((r, c), (r + self.rng.randrange(0, 3), c + self.rng.randrange(0, 2)))
        return r, c

    def a_insert_rc(self):
        at = self._select_some()
        if at is None:
            return "none"
        tables = self.view.tables
        which = self.rng.randrange(4)
        [lambda: tables.insert_rows(), lambda: tables.insert_rows(below=True),
         lambda: tables.insert_cols(), lambda: tables.insert_cols(right=True)][which]()
        return f"{['rows above', 'rows below', 'cols left', 'cols right'][which]} at {at}"

    def a_delete_rc(self):
        at = self._select_some()
        if at is None:
            return "none"
        tables = self.view.tables
        if self.rng.random() < 0.5:
            tables.delete_rows()
            return f"rows at {at}"
        tables.delete_cols()
        return f"cols at {at}"

    def _drag(self, start, end, steps=6):
        from PySide6.QtCore import QPoint, Qt
        from PySide6.QtTest import QTest
        port = self.view.viewport()
        QTest.mousePress(port, Qt.LeftButton, Qt.NoModifier, start)
        X.pump()
        for k in range(1, steps + 1):
            QTest.mouseMove(port, QPoint(start.x() + (end.x() - start.x()) * k // steps,
                                         start.y() + (end.y() - start.y()) * k // steps))
            X.pump()
        QTest.mouseRelease(port, Qt.LeftButton, Qt.NoModifier, end)

    def a_move_cells(self):
        from PySide6.QtCore import QPoint
        at = self._select_some()
        if at is None:
            return "none"
        run, tables = self.view.tables.item, self.view.tables
        t, l, b, r = tables.selection()
        box = run.block_rect(t, l, b, r)
        self.view.centerOn(run.mapToScene(box.center()))
        X.pump()
        edge = self.view.mapFromScene(run.mapToScene(box.topLeft() + _ptf(box.width() / 2, 0)))
        dx, dy = self.rng.randrange(-120, 160), self.rng.randrange(-80, 160)
        self._drag(QPoint(edge.x(), edge.y()), QPoint(edge.x() + dx, edge.y() + dy))
        return f"{(t, l, b, r)} by {dx},{dy}px"

    def a_fill(self):
        from PySide6.QtCore import QPoint
        at = self._select_some()
        if at is None:
            return "none"
        run = self.view.tables.item
        handle = run.fill_handle_rect()
        if handle is None:
            return "no handle"
        self.view.centerOn(run.mapToScene(handle.center()))
        X.pump()
        p = self.view.mapFromScene(run.mapToScene(handle.center()))
        if self.rng.random() < 0.3:
            from PySide6.QtCore import Qt
            from PySide6.QtTest import QTest
            QTest.mouseDClick(self.view.viewport(), Qt.LeftButton, Qt.NoModifier, QPoint(p.x(), p.y()))
            return "double-click"
        dy, dx = self.rng.randrange(0, 200), self.rng.randrange(0, 60) if self.rng.random() < 0.3 else 0
        self._drag(QPoint(p.x(), p.y()), QPoint(p.x() + dx, p.y() + dy))
        return f"drag {dx},{dy}"

    def a_copy_paste(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        at = self._select_some()
        if at is None:
            return "none"
        key = Qt.Key_X if self.rng.random() < 0.3 else Qt.Key_C
        QTest.keyClick(self.view.viewport(), key, Qt.ControlModifier)
        X.pump()
        run = self._run_for_cells()
        r, c = self.rng.randrange(0, 30), self.rng.randrange(0, 8)
        self._click(self._cell_point(run, r, c))
        X.pump()
        QTest.keyClick(self.view.viewport(), Qt.Key_V, Qt.ControlModifier)
        return f"{'cut' if key == Qt.Key_X else 'copy'} {at} -> ({r},{c})"

    def a_rename_sheet(self):
        runs = self.runs()
        if not runs:
            return "none"
        run = self.pick(runs)
        new = self.pick(NAMES + [n.upper() for n in self.sheet_names()] + self.sheet_names())
        self.view.tables.rename_to(run, new)
        return f"{run.sheet.name} -> {new}"

    def a_new_sheet(self):
        if self._run_for_cells() is None:
            return "none"
        self.win.sheet_tabs.new_sheet()
        return "+"

    def a_delete_sheet(self):
        tabs = self.win.sheet_tabs
        if self._run_for_cells() is None or not tabs.entries:
            return "none"
        index = self.rng.randrange(len(tabs.entries))
        tabs.delete_sheet(index)
        return f"tab {index}"

    def a_define_name(self):
        at = self._select_some()
        if at is None:
            return "none"
        name = self.pick(["W_total", "Span", "x1", "Load_2", "Z", "Height"])
        self.view.tables.define_name(name)
        return f"{name} at {at}"

    def a_freeze(self):
        if self._select_some() is None:
            return "none"
        action = self.pick([self.win.act_freeze_panes, getattr(self.win, "act_unfreeze_panes", None),
                            getattr(self.win, "act_freeze_top_row", None)])
        if action is None:
            return "none"
        action.trigger()
        return action.text()

    def a_drag_area(self):
        from PySide6.QtCore import QPoint
        run = self._run_for_cells()
        if run is None or run.paging is None:
            return "none"
        xs, ys = run.edges()
        paging = run.paging
        rows = paging.print_rows
        k = self.rng.random()
        if k < 0.5 and rows:
            local = _ptf(xs[min(2, len(xs) - 1)], ys[min(rows[1] + 1, len(ys) - 1)])
            what = "area bottom"
        elif k < 0.8:
            local = _ptf(xs[min(paging.last_col + 1, len(xs) - 1)], ys[min(3, len(ys) - 1)])
            what = "area right"
        elif len(paging.slices) > 1:
            local = _ptf(xs[1], ys[paging.slices[1][0]])
            what = "break"
        else:
            return "no break"
        self.view.centerOn(run.mapToScene(local))
        X.pump()
        p = self.view.mapFromScene(run.mapToScene(local))
        port = self.view.viewport()
        dx, dy = (self.rng.randrange(-80, 120), 0) if "right" in what else (0, self.rng.randrange(-100, 140))
        if self.rng.random() < 0.3:
            dy = port.height() - p.y() + 30 if dy else dy           # past the edge: autoscroll
            dx = port.width() - p.x() + 30 if dx else dx
        self._drag(QPoint(p.x(), p.y()), QPoint(p.x() + dx, p.y() + dy))
        return f"{what} by {dx},{dy}"

    def a_heading_select(self):
        from PySide6.QtCore import QPoint
        run = self._run_for_cells()
        if run is None:
            return "none"
        r = self.rng.randrange(0, 20)
        p = self._cell_point(run, r, 0)
        headings = getattr(self.view, "sheet_headings", None)
        left = 4
        if headings is not None:
            for _run, kind, rect in headings._strips:
                if kind == "row":
                    left = rect.center().x()
        self._click(QPoint(left, p.y()))
        return f"row {r}"

    # markups
    def a_draw_markup(self):
        from PySide6.QtCore import QPoint, Qt
        from PySide6.QtTest import QTest
        tool = self.pick(["rect", "ellipse", "line", "arrow", "cloud", "polyline"])
        try:
            self.view.set_tool(tool)
        except Exception:  # noqa: BLE001
            self.win.select_tool(tool)
        X.pump()
        port = self.view.viewport()
        x, y = self.rng.randrange(60, port.width() - 150), self.rng.randrange(60, port.height() - 150)
        w, h = self.rng.randrange(20, 140), self.rng.randrange(20, 120)
        if tool == "polyline":
            for px, py in ((x, y), (x + w, y), (x + w, y + h)):
                QTest.mouseClick(port, Qt.LeftButton, Qt.NoModifier, QPoint(px, py))
                X.pump()
            QTest.mouseDClick(port, Qt.LeftButton, Qt.NoModifier, QPoint(x, y + h))
        else:
            self._drag(QPoint(x, y), QPoint(x + w, y + h))
        X.pump()
        QTest.keyClick(port, Qt.Key_Escape)
        self.win.select_tool("select")
        return f"{tool} at {x},{y} {w}x{h}{' on a sheet' if self.shown() else ''}"

    def _markups(self):
        from calcforge.items.base import MarkupItem
        from calcforge.items.sheetpage import SheetRunItem
        out = []
        for frame in getattr(self.win.scene, "frames", []) or []:
            for item in frame.markups():
                if isinstance(item, MarkupItem) and not isinstance(item, SheetRunItem) \
                        and type(item).__name__ not in ("TableItem", "CalcItem"):
                    out.append(item)
        return out

    def _markup_point(self, item):
        from PySide6.QtCore import QPoint
        rect = item.mapRectToScene(item.local_rect()) if hasattr(item, "local_rect") else item.sceneBoundingRect()
        spot = rect.center()
        self.view.centerOn(spot)
        X.pump()
        p = self.view.mapFromScene(spot)
        return QPoint(p.x(), p.y())

    def a_move_markup(self):
        from PySide6.QtCore import QPoint
        items = [m for m in self._markups() if m.isVisible()]
        if not items:
            return "none"
        item = self.pick(items)
        self.win.select_tool("select")
        p = self._markup_point(item)
        dx, dy = self.rng.randrange(-100, 100), self.rng.randrange(-100, 100)
        self._drag(p, QPoint(p.x() + dx, p.y() + dy))
        return f"{type(item).__name__} by {dx},{dy}"

    def a_delete_markup(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        items = [m for m in self._markups() if m.isVisible()]
        if not items:
            return "none"
        item = self.pick(items)
        self.win.select_tool("select")
        self.view.scene().clearSelection()
        item.setSelected(True)
        self.win.refresh_selection()
        QTest.keyClick(self.view.viewport(), Qt.Key_Delete)
        return type(item).__name__

    def a_equation(self):
        from calcforge.ui import sheetview
        from calcforge.sheet.docbook import sheet_ident
        if self.shown() is not None:
            sheetview.leave(self.win)
            X.pump()
        doc = self.win.document
        plain = [i for i, p in enumerate(doc.pages) if p.sheet is None]
        if not plain:
            return "no plain page"
        index = self.pick(plain)
        self.win.go_to_page(index)
        X.pump()
        frame = doc.pages[index].frame
        if frame is None:
            return "no frame"
        names = [sheet_ident(n) for n in self.sheet_names()]
        ref = f"{self.pick(names)}.A{self.rng.randrange(1, 9)}" if names else "2"
        text = self.pick([f"v{self.rng.randrange(9)}:{ref}+1", f"w:{ref}*2'm", "q:3", f"{ref}=",
                          "r:q*v1", f"s:Z", "t:sum(1,2)", f"u:{ref}'kN"])
        from PySide6.QtCore import QPointF, Qt
        calc = self.view.calc
        calc.leave()
        self.win.select_tool("select")
        self.win.toggle_calc_mode(True)        # equations are typed in Calc mode
        self.view.scene().clearSelection()
        point = QPointF(self.rng.randrange(40, 400), self.rng.randrange(40, 700))
        spot = frame.mapToScene(point)
        self.view.centerOn(spot)
        X.pump()
        self._click(self.view.mapFromScene(spot))  # the red cross, where clicked
        X.pump()
        from PySide6.QtGui import QKeyEvent
        from PySide6.QtCore import QEvent
        from PySide6.QtWidgets import QApplication
        for ch in text:
            target = self.view.viewport()
            QApplication.sendEvent(target, QKeyEvent(QEvent.KeyPress, 0, Qt.NoModifier, ch))
            X.pump()
        calc.leave()
        if self.rng.random() < 0.5:
            self.win.toggle_calc_mode(False)
        return f"page {index}: {text}"

    def a_calc_mode(self):
        on = self.view.calc.mode != "calc"
        self.win.toggle_calc_mode(on)
        return "on" if on else "off"

    def _blank_spot(self):
        """A point in the viewport over the page in view."""
        port = self.view.viewport()
        return _pt(self.rng.randrange(60, max(61, port.width() - 200)),
                   self.rng.randrange(60, max(61, port.height() - 150)))

    def a_click_tool(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        tool = self.pick(["stamp", "note", "flag", "count", "checkbox", "stamp"])
        self.win.select_tool(tool)
        X.pump()
        spot = self._blank_spot()
        self._click(spot)
        X.pump()
        if tool == "count" and self.rng.random() < 0.6:
            self._click(spot + _pt(40, 10))
            X.pump()
        if tool in ("note", "checkbox") and self.rng.random() < 0.6:
            QTest.keyClicks(self.view.viewport(), self.pick(TEXTS))
            X.pump()
        QTest.keyClick(self.view.viewport(), Qt.Key_Escape)
        self.win.select_tool("select")
        return f"{tool} at {spot.x()},{spot.y()}{' on a sheet' if self.shown() else ''}"

    def a_text_box(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        tool = self.pick(["text", "text", "callout"])
        self.win.select_tool(tool)
        X.pump()
        spot = self._blank_spot()
        if tool == "callout":
            self._click(spot)                       # where the arrow points
            X.pump()
            self._drag(spot + _pt(60, -60), spot + _pt(180, -20))
        else:
            self._drag(spot, spot + _pt(self.rng.randrange(20, 160), self.rng.randrange(15, 80)))
        X.pump()
        words = self.pick(TEXTS + ["=Z*2", "Span 6 m", "Check"])
        QTest.keyClicks(self.view.viewport(), words)
        X.pump()
        how = self.rng.random()
        if how < 0.4:
            QTest.keyClick(self.view.viewport(), Qt.Key_Escape)
        elif how < 0.8:
            self.win.select_tool("select")
            self._click(self._blank_spot())          # clicked away
        self.win.select_tool("select")
        return f"{tool}: {words!r} {'Esc' if how < 0.4 else 'click away' if how < 0.8 else 'left typing'}"

    def a_markup_keys(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        items = [m for m in self._markups() if m.isVisible()]
        if not items:
            return "none"
        item = self.pick(items)
        self.win.select_tool("select")
        if self.rng.random() < 0.5:
            self._click(self._markup_point(item))    # clicked on
            how = "clicked"
        else:
            self.view.scene().clearSelection()       # as the Markups list picks it
            item.setSelected(True)
            self.win.refresh_selection()
            how = "listed"
        X.pump()
        port = self.view.viewport()
        keys = self.pick([
            [(Qt.Key_Delete, Qt.NoModifier)],
            [(Qt.Key_Right, Qt.NoModifier)] * 3,
            [(Qt.Key_Up, Qt.ShiftModifier)],
            [(Qt.Key_C, Qt.ControlModifier), (Qt.Key_V, Qt.ControlModifier)],
            [(Qt.Key_X, Qt.ControlModifier), (Qt.Key_V, Qt.ControlModifier)],
            [(Qt.Key_D, Qt.ControlModifier)],
            [(Qt.Key_A, Qt.ControlModifier), (Qt.Key_Right, Qt.NoModifier)],
            [(Qt.Key_Escape, Qt.NoModifier)],
        ])
        for key, mods in keys:
            QTest.keyClick(port, key, mods)
            X.pump()
        names = {Qt.Key_Delete: "Del", Qt.Key_Right: "→", Qt.Key_Up: "↑", Qt.Key_C: "C",
                 Qt.Key_V: "V", Qt.Key_X: "X", Qt.Key_D: "D", Qt.Key_A: "A", Qt.Key_Escape: "Esc"}
        return f"{type(item).__name__} {how}: " + " ".join(
            ("Ctrl+" if m == Qt.ControlModifier else "Shift+" if m == Qt.ShiftModifier else "")
            + names[k] for k, m in keys)

    def a_resize_markup(self):
        items = [m for m in self._markups() if m.isVisible() and hasattr(m, "handle_points")]
        if not items:
            return "none"
        item = self.pick(items)
        self.win.select_tool("select")
        self._click(self._markup_point(item))
        X.pump()
        try:
            handles = item.handle_points()
        except Exception:  # noqa: BLE001
            return "no handles"
        if not handles:
            return "no handles"
        name = self.pick(sorted(handles))
        spot = item.mapToScene(handles[name])
        self.view.centerOn(spot)
        X.pump()
        p = self.view.mapFromScene(spot)
        dx, dy = self.rng.randrange(-60, 80), self.rng.randrange(-60, 80)
        self._drag(p, p + _pt(dx, dy))
        return f"{type(item).__name__} {name} by {dx},{dy}"

    def a_edit_equation(self):
        from PySide6.QtCore import QEvent, Qt
        from PySide6.QtGui import QKeyEvent
        from PySide6.QtTest import QTest
        from PySide6.QtWidgets import QApplication
        regions = [m for f in getattr(self.win.scene, "frames", []) or [] for m in f.markups()
                   if type(m).__name__ == "CalcItem" and m.isVisible()]
        if not regions:
            return "none"
        region = self.pick(regions)
        box = region.local_rect() if hasattr(region, "local_rect") else region.boundingRect()
        spot = region.mapToScene(box.center())
        self.view.centerOn(spot)
        X.pump()
        p = self.view.mapFromScene(spot)
        QTest.mouseDClick(self.view.viewport(), Qt.LeftButton, Qt.NoModifier, p)
        X.pump()
        typed = self.pick(["+1", "*2", "\b\b", "'m", "^2"])
        for ch in typed.replace("\b", "\x08"):
            if ch == "\x08":
                QTest.keyClick(self.view.viewport(), Qt.Key_Backspace)
            else:
                QApplication.sendEvent(self.view.viewport(), QKeyEvent(QEvent.KeyPress, 0, Qt.NoModifier, ch))
            X.pump()
        if self.rng.random() < 0.5:
            QTest.keyClick(self.view.viewport(), Qt.Key_Escape)
        else:
            self._click(self._blank_spot())
        self.view.calc.leave()
        return f"{typed!r}"

    def a_style(self):
        items = [m for m in self._markups() if m.isVisible()]
        if not items:
            return "none"
        picked = [self.pick(items) for _ in range(self.rng.randrange(1, 3))]
        self.view.scene().clearSelection()
        for item in picked:
            item.setSelected(True)
        self.win.refresh_selection()
        X.pump()
        win = self.win
        change = self.pick([
            ("line colour", lambda: win._style_stroke(self.pick(["#e03131", "#1c7ed6", "#000000"]))),
            ("fill", lambda: win._style_fill(self.pick(["#ffd43b", "#ffffff", ""]))),
            ("width", lambda: win._style_width(self.pick([0.5, 1.5, 4.0]))),
            ("opacity", lambda: win._style_opacity(self.pick([20, 60, 100]))),
            ("dash", lambda: win._style_dash(self.pick(["dashed 1", "dashed 4", "solid"]))),
            ("font", lambda: win._style_font(self.pick([8.0, 12.0, 20.0]))),
            ("hatch", lambda: win._style_hatch(self.pick(["", "brick", "up"]))),
        ])
        change[1]()
        return f"{change[0]} of {', '.join(type(i).__name__ for i in picked)}"

    def a_select_all_delete(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        if self.shown() is not None and self.rng.random() < 0.5:
            from calcforge.ui import sheetview
            sheetview.leave(self.win)
            X.pump()
        self.win.select_tool("select")
        port = self.view.viewport()
        QTest.keyClick(port, Qt.Key_Escape)
        QTest.keyClick(port, Qt.Key_A, Qt.ControlModifier)
        X.pump()
        picked = len(self.view.scene().selectedItems())
        QTest.keyClick(port, Qt.Key_Delete)
        return f"{picked} picked"

    SKIP_ENTRIES = ("print", "export", "save", "email", "send", "open", "properties",
                    "keyboard", "customize", "add to tool chest", "set as default", "flatten")

    def a_context_entry(self):
        from PySide6.QtWidgets import QMenu
        where = self.rng.random()
        menu = None
        what = ""
        if where < 0.35 and self.view.tables.item is not None:
            menu, what = self.view.tables.cell_menu(), "cell"
        elif where < 0.65:
            items = [m for m in self._markups() if m.isVisible()]
            if items:
                item = self.pick(items)
                self.view.scene().clearSelection()
                item.setSelected(True)
                self.win.refresh_selection()
                spot = item.mapToScene(item.boundingRect().center())
                menu, what = self.win.build_context_menu(item, spot), type(item).__name__
        elif where < 0.85:
            index = self.rng.randrange(len(self.win.document.pages))
            menu, what = self.win.page_menu(index), f"thumbnail {index}"
        if menu is None:
            frame = self.win.document.pages[self.win.current_index].frame
            if frame is None:
                return "none"
            spot = frame.mapToScene(frame.page_rect().center())
            menu, what = self.win.build_context_menu(None, spot), "page"
        entries = []

        def walk(m, path):
            m.aboutToShow.emit()
            for a in m.actions():
                label = a.text().replace("&", "")
                if a.menu() is not None:
                    walk(a.menu(), path + [label])
                elif label and a.isEnabled() and not a.isSeparator() \
                        and not any(k in label.lower() for k in self.SKIP_ENTRIES):
                    entries.append((path + [label], a))
        walk(menu, [])
        if not entries:
            return f"{what}: nothing"
        path, action = self.pick(entries)
        action.trigger()
        X.pump()
        if isinstance(menu, QMenu):
            menu.deleteLater()
        return f"{what}: {' ▸ '.join(path)}"

    def a_format_cells(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        at = self._select_some()
        if at is None:
            return "none"
        key = self.pick([Qt.Key_B, Qt.Key_I, Qt.Key_U, Qt.Key_1, Qt.Key_5])
        mods = Qt.ControlModifier | (Qt.ShiftModifier if key in (Qt.Key_1, Qt.Key_5) else Qt.NoModifier)
        QTest.keyClick(self.view.viewport(), key, mods)
        return f"{at} Ctrl+{'Shift+' if mods & Qt.ShiftModifier else ''}{chr(key.value)}"

    # the rest
    def a_undo(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        n = self.rng.randrange(1, 4)
        for _ in range(n):
            if self.rng.random() < 0.5:
                QTest.keyClick(self.view.viewport(), Qt.Key_Z, Qt.ControlModifier)
            else:
                self.win.undo_stack.undo()
            X.pump()
        return f"x{n}"

    def a_redo(self):
        n = self.rng.randrange(1, 3)
        for _ in range(n):
            self.win.undo_stack.redo()
            X.pump()
        return f"x{n}"

    def a_zoom(self):
        from PySide6.QtCore import QPoint, QPointF, Qt
        from PySide6.QtGui import QWheelEvent
        from PySide6.QtWidgets import QApplication
        port = self.view.viewport()
        at = QPointF(self.rng.randrange(20, port.width() - 20), self.rng.randrange(20, port.height() - 20))
        notches = self.rng.choice([-2, -1, 1, 2, 3])
        for _ in range(abs(notches)):
            QApplication.sendEvent(port, QWheelEvent(at, QPointF(port.mapToGlobal(at.toPoint())), QPoint(0, 0),
                                                     QPoint(0, 120 if notches > 0 else -120), Qt.NoButton,
                                                     Qt.ControlModifier, Qt.NoScrollPhase, False))
            X.pump()
        getattr(self.view, "finish_scrolling", lambda: None)()
        return f"{notches} notches"

    def a_scroll(self):
        bar = self.view.verticalScrollBar() if self.rng.random() < 0.7 else self.view.horizontalScrollBar()
        span = max(bar.maximum() - bar.minimum(), 1)
        bar.setValue(bar.minimum() + self.rng.randrange(0, span + 1))
        return f"{bar.orientation()} to {bar.value()}"

    SKIP = ("exit", "quit", "close", "new window", "open", "print", "save", "export", "import",
            "recent", "about", "help", "manual", "preferences", "full screen", "batch",
            "sign", "flatten", "compare", "email", "send", "check for", "update", "reset",
            "split document", "extract", "combine", "insert pages", "pdf", "keyboard",
            "customize", "workspace", "plugin", "tool chest", "profile", "theme")

    def a_menu_command(self):
        actions = []

        def walk(menu):
            for a in menu.actions():
                if a.menu() is not None:
                    walk(a.menu())
                elif a.text() and a.isEnabled() and not a.isSeparator():
                    actions.append(a)
        for top in self.win.menuBar().actions():
            if top.menu() is not None:
                walk(top.menu())
        actions = [a for a in actions if not any(s in a.text().lower().replace("&", "") for s in self.SKIP)]
        if not actions:
            return "none"
        action = self.pick(actions)
        action.trigger()
        return action.text().replace("&", "")

    # -- the checks --------------------------------------------------------------------
    def invariants(self) -> list:
        problems = []
        win, doc = self.win, self.win.document
        frames = getattr(win.scene, "frames", []) or []
        if len(frames) != len(doc.pages):
            problems.append(f"{len(frames)} frames for {len(doc.pages)} pages")
        strip = win.pages_panel.list
        if strip.count() != len(doc.pages):
            problems.append(f"{strip.count()} thumbnails for {len(doc.pages)} pages")
        uids = {p.sheet for p in doc.pages if p.sheet is not None}
        found = {r.uid for r in self.runs()}
        missing = [u for u in uids if u not in found and _run_waiting(doc, u) is False]
        if missing:
            problems.append(f"sheets with pages but no cells: {missing[:3]}")
        # every run's pages together, its cells on the first
        for uid in uids:
            idx = [i for i, p in enumerate(doc.pages) if p.sheet == uid]
            if idx != list(range(idx[0], idx[-1] + 1)):
                problems.append(f"sheet {uid[:6]}'s pages apart: {idx}")
        from calcforge.sheet import formula as F
        from calcforge.sheet.docbook import book_for
        wb = book_for(doc).workbook
        for sheet in wb.sheets:
            for (r, c), cell in list(sheet.cells.items()):
                if cell.is_formula:
                    try:
                        F.tokenize(cell.input[1:])
                    except F.FormulaError as e:
                        if not _typed_badly(cell.input):
                            problems.append(f"unreadable formula {sheet.name}!({r},{c}) {cell.input!r}: {e}")
        for dn in wb.names.values():
            try:
                F.tokenize(dn.refers_to)
            except F.FormulaError as e:
                problems.append(f"unreadable name {dn.name} = {dn.refers_to!r}: {e}")
        names = [s.name.lower() for s in wb.sheets]
        if len(names) != len(set(names)):
            problems.append(f"two sheets of one name: {sorted(n for n in names if names.count(n) > 1)}")
        return problems

    def signature(self) -> str:
        draft = getattr(self.view, "_draft", None)
        if draft is None or draft.scene() is None:
            return X.doc_signature(self.win)
        # a markup half drawn (a measurement waiting for its second click) is
        # not the document's yet: an undo is not expected to bring it back
        try:
            data = self.win.document.to_dict()
        except Exception as error:  # noqa: BLE001
            return f"<to_dict raised {error!r}>"
        uid = getattr(draft, "uid", None)
        for page in data.get("pages", []):
            page["items"] = [it for it in page.get("items") or []
                             if not (isinstance(it, dict) and it.get("uid") == uid)]
        return json.dumps(X._normal(data), sort_keys=True, default=str)

    # -- the run -----------------------------------------------------------------------
    def settle(self) -> None:
        """As Undo History settles the canvas before it undoes: the words
        being typed end (joining the step that put their box down), a markup
        half drawn goes. The checks below undo the stack itself."""
        self.view.end_item_edit()
        if self.view._draft is not None:
            self.view.cancel_draft()
        X.settle(5)

    def step(self, i: int, check_undo: bool) -> None:
        name, fn = self.choose()
        stack = self.win.undo_stack
        if check_undo:
            self.settle()
        before_sig = self.signature() if check_undo else None
        before_index = stack.index()
        X.CURRENT.clear()
        X.CURRENT.update(errors=[], qt_messages=[])
        started = time.perf_counter()
        desc = ""
        try:
            desc = fn()
        except Exception:  # noqa: BLE001
            X.CURRENT["errors"].append(traceback.format_exc())
        X.settle(15)
        took = time.perf_counter() - started
        entry = {"step": i, "action": name, "detail": desc, "took": round(took, 2)}
        self.log.append(entry)
        if self.verbose:
            print(f"  {i:4d} {name}: {desc} ({took:.2f}s)", flush=True)
        if self.collect:
            # what the step let go of, looked over now: a crash in it is this step's
            import gc
            gc.collect()
        errors = list(X.CURRENT.get("errors") or [])
        qt_bad = [m for m in X.CURRENT.get("qt_messages") or [] if m.startswith(("critical", "fatal"))]
        if errors or qt_bad:
            self.fail(i, "exception", errors + qt_bad)
            return
        problems = self.invariants()
        if problems:
            self.fail(i, "broken document", problems)
            return
        if took > 4.0:
            self.fail(i, "slow", [f"{name} took {took:.1f}s"])
        if check_undo and name not in ("undo", "redo"):
            self.settle()
            done = stack.index() - before_index
            if done > 0 and stack.index() == stack.count():
                after_sig = self.signature()
                for _ in range(done):
                    stack.undo()
                X.settle(15)
                back = self.signature()
                if back != before_sig:
                    self.fail(i, "undo did not put it back", [_diff(before_sig, back)])
                for _ in range(done):
                    stack.redo()
                X.settle(15)
                again = self.signature()
                if again != after_sig:
                    self.fail(i, "redo did not give it again", [_diff(after_sig, again)])
                errors = list(X.CURRENT.get("errors") or [])
                if errors:
                    self.fail(i, "exception in undo/redo", errors)

    def round_trip(self, i: int) -> None:
        from calcforge.ui import sheetview
        if sheetview.run_shown(self.win.scene) is not None:
            sheetview.leave(self.win)
            X.pump()
        path = os.path.join(self.tmp, f"doc-{i}.pdf")
        X.CURRENT.clear()
        X.CURRENT.update(errors=[], qt_messages=[])
        self.win.document.path = path
        try:
            ok = self.win.save_document()
        except Exception:  # noqa: BLE001
            self.fail(i, "save raised", [traceback.format_exc()])
            return
        if not ok:
            self.fail(i, "save failed", [])
            return
        mine = _comparable(self.win.document.to_dict())
        other = X.new_window()
        other.interactive_prompts = False
        try:
            other.open_path(path)
            X.settle(40)
            theirs = _comparable(other.document.to_dict())
            if mine != theirs:
                self.fail(i, "save and open gave another document",
                          [_diff(json.dumps(mine, sort_keys=True, default=str),
                                 json.dumps(theirs, sort_keys=True, default=str))])
            errors = list(X.CURRENT.get("errors") or [])
            if errors:
                self.fail(i, "exception opening the saved document", errors)
        finally:
            X.dispose(other)

    def fail(self, i: int, kind: str, details: list) -> None:
        record = {"seed": self.seed, "step": i, "kind": kind, "details": details[:4],
                  "last_steps": self.log[-12:]}
        self.failures.append(record)
        print(f"FAIL seed {self.seed} step {i}: {kind}\n" +
              "\n".join(str(d)[:3000] for d in details[:2]), flush=True)
        with open(os.path.join(self.out, f"seed{self.seed}.jsonl"), "a") as fh:
            fh.write(json.dumps(record, default=str) + "\n")

    def run(self, steps: int, undo_every: float = 0.35, save_every: int = 120,
            stop_after: int = 6) -> list:
        for i in range(steps):
            self.step(i, check_undo=self.rng.random() < undo_every)
            if save_every and i and i % save_every == 0:
                self.round_trip(i)
            if len(self.failures) >= stop_after:
                break
            if any(f["kind"] in ("exception", "broken document") for f in self.failures[-1:]) \
                    and self.failures[-1]["step"] == i:
                # the window may be in a bad way: a fresh document, the run goes on
                try:
                    self.win.new_document(confirm=False)
                    X.pump(3)
                except Exception:  # noqa: BLE001
                    break
        return self.failures


def _pt(x, y):
    from PySide6.QtCore import QPoint
    return QPoint(int(x), int(y))


def _ptf(x, y):
    from PySide6.QtCore import QPointF
    return QPointF(float(x), float(y))


def _run_waiting(doc, uid) -> bool:
    """Whether sheet *uid*'s cells are still a record waiting for its page to
    be drawn (not a fault)."""
    for page in doc.pages:
        if page.sheet == uid:
            for record in getattr(page, "_pending_items", None) or ():
                if isinstance(record, dict) and record.get("uid") == uid:
                    return True
            return page.frame is None
    return False


def _typed_badly(text: str) -> bool:
    """A formula the reader typed that doesn't read is kept as typed (Excel
    refuses it; CalcForge keeps it in red): not a fault of the app."""
    return text.rstrip().endswith(("(", "+", "=", ",")) or text.strip() in ("=",)


#: what a save and open changes on purpose: a page opened from the saved PDF
#: is drawn on its own page of that PDF (its pdf_key, pdf_page_index)
VOLATILE_KEYS = {"path", "modified", "revision", "doc_id", "saved_at", "file_time",
                 "pdf_key", "pdf_page_index", "written_here"}


def _comparable(data):
    if isinstance(data, dict):
        return {k: _comparable(v) for k, v in sorted(data.items())
                if k not in VOLATILE_KEYS and k not in X.VOLATILE}
    if isinstance(data, (list, tuple)):
        return [_comparable(v) for v in data]
    if isinstance(data, float):
        return round(data, 2)
    if isinstance(data, bytes):
        return f"<{len(data)} bytes>"
    return data


def _diff(a: str, b: str, width: int = 260) -> str:
    """Where two signatures part, with a little either side."""
    k = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), min(len(a), len(b)))
    return f"at {k}: …{a[max(0, k - width):k + width]}…\n  vs …{b[max(0, k - width):k + width]}…"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="cmd", required=True)
    run = sub.add_parser("run")
    run.add_argument("--seeds", default="1-4")
    run.add_argument("--steps", type=int, default=300)
    run.add_argument("--out", default=os.path.join(tempfile.gettempdir(), "calcforge-monkey"))
    run.add_argument("--verbose", action="store_true")
    rep = sub.add_parser("replay")
    rep.add_argument("--seed", type=int, required=True)
    rep.add_argument("--upto", type=int, required=True)
    rep.add_argument("--out", default=os.path.join(tempfile.gettempdir(), "calcforge-monkey"))
    rep.add_argument("--verbose", action="store_true")
    args = parser.parse_args(argv)
    os.makedirs(args.out, exist_ok=True)
    if args.cmd == "run":
        lo, _, hi = args.seeds.partition("-")
        seeds = range(int(lo), int(hi or lo) + 1)
        total = []
        for seed in seeds:
            print(f"seed {seed}", flush=True)
            session = Session(seed, args.out, args.verbose)
            total += session.run(args.steps)
            X.dispose(session.win)
        print(f"MONKEY_DONE failures={len(total)}", flush=True)
        return 1 if total else 0
    session = Session(args.seed, args.out, args.verbose)
    session.run(args.upto + 1, stop_after=10 ** 6)
    print(f"MONKEY_DONE failures={len(session.failures)}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
