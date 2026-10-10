"""Opening Excel workbooks (spreadsheet phase 6): each worksheet a run of
spreadsheet pages with its cells, formulas, formats, sizes, merges, names,
comments, validation, conditional formatting and page layout; Excel tables'
references as cells; charts and pictures left out and said so; nothing
written back to Excel."""
from __future__ import annotations

import datetime

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table  # noqa: E402

from calcforge.items.sheetpage import SheetRunItem  # noqa: E402
from calcforge.sheet.pagination import options  # noqa: E402
from calcforge.sheet.xlsx import read_workbook  # noqa: E402


def a_workbook(path, with_chart=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Beams"
    rows = [("Member", "Load", "Span", "Moment"), ("B1", 10, 4, None), ("B2", 20, 5, None),
            ("B3", 30, 6, None)]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            if v is not None:
                ws.cell(r, c, v)
    for r in (2, 3, 4):
        ws.cell(r, 4, f"=B{r}*C{r}^2/8")
    ws.add_table(Table(displayName="BeamTbl", ref="A1:D4"))
    ws["F1"], ws["G1"] = "Total", "=SUM(BeamTbl[Moment])"
    ws["F2"], ws["G2"] = "Heaviest", "=_xlfn.XLOOKUP(MAX(B2:B4),B2:B4,A2:A4)"
    ws["F3"], ws["G3"] = "Date", datetime.date(2026, 10, 10)
    ws["G3"].number_format = "d/m/yyyy"
    ws["F4"], ws["G4"] = "Factored", "='Load Cases'!B2*G1"
    ws["H2"] = "=BeamTbl[@Load]*2"
    ws["F5"], ws["G5"] = "Text", "=not a formula"
    ws["G5"].data_type = "s"                     # text that starts with =
    ws["A1"].font = Font(bold=True, color="FFFFFF", name="Arial", sz=12)
    ws["A1"].fill = PatternFill("solid", fgColor="4472C4")
    ws["B2"].border = Border(bottom=Side("medium", color="FF0000"))
    ws["D2"].number_format = '0.0 "kN·m"'
    ws["D2"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 15
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["E"].hidden = True
    ws.merge_cells("F7:G7")
    ws["F7"] = "merged"
    dv = DataValidation(type="list", formula1='"B1,B2,B3"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A10")
    ws.conditional_formatting.add("D2:D4", CellIsRule(
        operator="greaterThan", formula=["50"], font=Font(color="9C0006"),
        fill=PatternFill("solid", bgColor="FFC7CE")))
    ws.conditional_formatting.add("B2:B4", ColorScaleRule(
        start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    ws["A2"].comment = Comment("first beam", "me")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=30))
    wb.defined_names["Loads"] = DefinedName("Loads", attr_text="Beams!$B$2:$B$4")
    wb.defined_names["gee"] = DefinedName("gee", attr_text="9.81")
    ws2 = wb.create_sheet("Load Cases")
    ws2["A1"], ws2["B1"], ws2["A2"], ws2["B2"] = "Case", "Factor", "ULS", 1.5
    ws2["C2"] = "=SEQUENCE(3)"
    ws2["D2"] = "=SUM(Loads)*gee"
    ws2.sheet_state = "hidden"
    if with_chart:
        from openpyxl.chart import BarChart, Reference
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        ws.add_chart(chart, "J2")
    wb.save(path)
    return path


def test_the_reader_brings_cells_formats_and_layout(tmp_path):
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    got = read_workbook(path)
    beams, cases = got.sheets
    assert (beams.name, cases.name) == ("Beams", "Load Cases")
    assert cases.hidden and not beams.hidden
    cells = {(r, c): (text, beams.data["styles"][s]) for r, c, text, s, *_ in beams.data["cells"]}
    assert cells[(1, 3)][0] == "=B2*C2^2/8"
    assert cells[(0, 6)][0] == "=SUM($D$2:$D$4)", "an Excel table's column as its cells"
    assert cells[(1, 7)][0] == "=$B2*2", "this row of an Excel table"
    assert cells[(1, 6)][0] == "=XLOOKUP(MAX(B2:B4),B2:B4,A2:A4)", "no _xlfn. prefixes"
    assert cells[(2, 6)][0] == "46305" and cells[(2, 6)][1]["number_format"] == "d/m/yyyy"
    assert cells[(4, 6)][0] == "'=not a formula"
    head = cells[(0, 0)][1]
    assert head["bold"] and head["color"] == "#FFFFFF" and head["fill"] == "#4472C4"
    assert head["font"] == "Arial" and head["size"] == 12
    assert cells[(1, 1)][1]["bottom"] == ["medium", "#FF0000"]
    assert cells[(1, 3)][1]["number_format"] == '0.0 "kN·m"'
    assert cells[(1, 3)][1]["h_align"] == "center" and cells[(1, 3)][1]["wrap"]
    comment = next(e for e in beams.data["cells"] if e[:2] == [1, 0])
    assert comment[4] == "first beam"
    assert beams.data["widths"]["0"] == pytest.approx(78.75)      # 15 characters
    assert beams.data["heights"]["0"] == 24
    assert beams.data["hidden_cols"] == [4]
    assert beams.data["merges"] == [[6, 5, 6, 6]]
    assert beams.paper == "A3" and beams.orientation == "landscape"
    page = beams.data["page"]
    assert page["titles"] == [0, 0] and page["center_h"] and page["breaks"] == [30]
    rules = beams.data["cond_rules"]
    assert rules[0]["type"] == "cell" and rules[0]["op"] == "greater" and rules[0]["a"] == "50"
    assert rules[0]["format"] == {"color": "#9C0006", "fill": "#FFC7CE"}
    assert rules[1]["type"] == "scale" and rules[1]["colors"] == ["#F8696B", "#63BE7B"]
    (dv,) = beams.data["validations"]
    assert dv["type"] == "list" and dv["source"] == "B1,B2,B3" and dv["ranges"] == [[9, 0, 9, 0]]
    names = {n[0]: n[1] for n in beams.data["names"]}
    assert names == {"Loads": "Beams!$B$2:$B$4", "gee": "9.81"}
    assert got.left_out == []


def test_a_worksheet_named_like_a_sheet_already_here_is_renamed(tmp_path):
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    got = read_workbook(path, taken={"Load Cases"})
    assert [s.name for s in got.sheets] == ["Beams", "Load Cases (2)"]
    cells = {(r, c): text for r, c, text, *_ in got.sheets[0].data["cells"]}
    assert cells[(3, 6)] == "='Load Cases (2)'!B2*G1", "the formulas follow the new name"


def test_only_the_worksheets_chosen_come_in(tmp_path):
    """A formula on a kept sheet reading one left out keeps Excel's last
    value, as Excel's Break Links does; a workbook name moves to a kept sheet
    unless it reads one left out."""
    from calcforge.sheet.xlsx import keep_only
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    got = read_workbook(path)
    got.sheets[0].excel_values[(3, 6)] = 326.25
    assert keep_only(got, ["Beams"]) == 1
    assert [s.name for s in got.sheets] == ["Beams"]
    cells = {(r, c): text for r, c, text, *_ in got.sheets[0].data["cells"]}
    assert cells[(3, 6)] == "326.25"
    assert cells[(0, 6)] == "=SUM($D$2:$D$4)", "formulas on this sheet are untouched"
    assert any("not brought in" in what for what in got.left_out)

    got = read_workbook(path)
    assert keep_only(got, ["load cases"]) == 1
    (cases,) = got.sheets
    assert {n[0] for n in cases.data["names"]} == {"gee"}, "Loads read Beams; gee is a constant"
    cells = {(r, c): text for r, c, text, *_ in cases.data["cells"]}
    assert cells[(1, 3)] == "", "it read the name of Beams' cells"
    assert cells[(1, 2)] == "=SEQUENCE(3)"


def test_insert_excel_asks_which_worksheets(w, tmp_path):
    from PySide6.QtWidgets import QApplication
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    asked = []
    w.choose_workbook_sheets = lambda names: asked.append(names) or ["Load Cases"]
    assert w.insert_workbook(path, 0)
    QApplication.processEvents()
    assert asked == [["Beams", "Load Cases"]]
    assert set(runs(w)) == {"Load Cases"}
    w.choose_workbook_sheets = lambda names: []           # cancelled
    pages = len(w.document.pages)
    assert not w.insert_workbook(path, 0)
    assert len(w.document.pages) == pages


def test_the_sheet_chooser_warns_of_a_sheet_read_but_left_out(qapp, tmp_path, monkeypatch):
    """Beams reads Load Cases ('Load Cases'!B2); Load Cases reads Beams through
    the name Loads. Leaving either out is said at once, and on OK the reader
    can tick it too."""
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QMessageBox
    from calcforge.sheet.xlsx import sheet_links
    from calcforge.ui.xlsximport import SheetChooser
    got = read_workbook(a_workbook(str(tmp_path / "beams.xlsx")))
    assert sheet_links(got) == {"Beams": {"Load Cases"}, "Load Cases": {"Beams"}}
    chooser = SheetChooser(None, got)
    assert chooser.warning.isHidden()
    chooser.list.item(1).setCheckState(Qt.Unchecked)
    assert chooser.missing() == [("Beams", "Load Cases")]
    assert "“Beams” reads “Load Cases”" in chooser.warning.text()
    assert not chooser.warning.isHidden()

    def answer(label):
        def exec_(box):
            box._clicked = next(b for b in box.buttons() if b.text() == label)
            return 0
        return exec_
    monkeypatch.setattr(QMessageBox, "clickedButton", lambda box: box._clicked)
    monkeypatch.setattr(QMessageBox, "exec", answer("Tick those too"))
    chooser._accept()
    assert chooser.chosen() == ["Beams", "Load Cases"] and chooser.warning.isHidden()
    chooser.list.item(0).setCheckState(Qt.Unchecked)
    monkeypatch.setattr(QMessageBox, "exec", answer("Bring in as chosen"))
    chooser._accept()
    assert chooser.chosen() == ["Load Cases"]


def test_the_sheet_chooser_ticks_them_all_to_start(qapp, tmp_path):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QDialogButtonBox
    from calcforge.ui.xlsximport import SheetChooser
    got = read_workbook(a_workbook(str(tmp_path / "beams.xlsx")))
    chooser = SheetChooser(None, got)
    assert chooser.chosen() == ["Beams", "Load Cases"]
    assert "hidden in Excel" in chooser.list.item(1).text()
    chooser.list.item(0).setCheckState(Qt.Unchecked)
    assert chooser.chosen() == ["Load Cases"]
    chooser._tick_all(Qt.Unchecked)
    assert not chooser.buttons.button(QDialogButtonBox.Ok).isEnabled()
    chooser._tick_all(Qt.Checked)
    assert chooser.buttons.button(QDialogButtonBox.Ok).isEnabled()


def test_excels_print_scaling_comes_across(tmp_path):
    """Adjust to N%, and Fit to page — 1 by 1 when the file names neither."""
    path = str(tmp_path / "scaled.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjusted"
    ws["A1"] = 1
    ws.page_setup.scale = 85
    fit = wb.create_sheet("Fitted")
    fit["A1"] = 1
    fit.sheet_properties.pageSetUpPr.fitToPage = True
    fit.print_title_cols = "A:B"
    fit.print_title_rows = "1:2"
    wb.save(path)
    got = read_workbook(path)
    adjusted, fitted = (s.data["page"] for s in got.sheets)
    assert adjusted["scale"] == 85 and not adjusted.get("fit_width")
    assert fitted["fit_width"] and fitted["fit_tall"] == 1
    assert fitted["title_cols"] == [0, 1] and fitted["titles"] == [0, 1]
    assert not any("scale" in what for what in got.left_out)


def test_excels_newer_validations_come_across_without_a_warning(tmp_path):
    """A list from another sheet is written in Excel's 2010 extension, which
    openpyxl drops with a warning on the console; it is read here instead."""
    import shutil
    import warnings
    import zipfile
    path = str(tmp_path / "lists.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Input"
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(["C25", "C32", "C40"], 1):
        lists.cell(i, 1, v)
    wb.save(path)
    ext = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
           'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
           '<x14:dataValidation type="list" allowBlank="1" showErrorMessage="1" '
           'errorTitle="Grade" error="Pick a grade"><x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f>'
           '</x14:formula1><xm:sqref>B2:B10</xm:sqref></x14:dataValidation></x14:dataValidations>'
           '</ext></extLst>')
    patched = str(tmp_path / "patched.xlsx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(patched, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.decode().replace("</worksheet>", ext + "</worksheet>").encode()
            out.writestr(item, data)
    shutil.copy(patched, path)
    with warnings.catch_warnings():
        warnings.simplefilter("error")                 # nothing said on the console
        got = read_workbook(path)
    (rule,) = got.sheets[0].data["validations"]
    assert rule["type"] == "list" and rule["source"] == "=Lists!$A$1:$A$3"
    assert rule["ranges"] == [[1, 1, 9, 1]] and rule["blank"] and rule["error"] == "Pick a grade"


def test_charts_and_pictures_are_left_out_and_said(tmp_path):
    path = a_workbook(str(tmp_path / "chart.xlsx"), with_chart=True)
    got = read_workbook(path)
    assert any("chart" in what for what in got.left_out)


@pytest.fixture
def w(window):
    window.show()
    window.view.set_zoom(1.0)
    return window


def runs(w):
    return {i.sheet.name: i for i in w.view.scene().items() if isinstance(i, SheetRunItem)}


def value(run, row, col):
    cell = run.sheet.cells.get((row, col))
    return None if cell is None else cell.value


def test_opening_an_xlsx_gives_its_worksheets_as_spreadsheet_pages(w, tmp_path):
    from PySide6.QtWidgets import QApplication
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    w.interactive_prompts = False
    w.open_from_command_line(path)
    QApplication.processEvents()
    QApplication.processEvents()
    got = runs(w)
    assert set(got) == {"Beams", "Load Cases"}
    beams, cases = got["Beams"], got["Load Cases"]
    assert [p.sheet is not None for p in w.document.pages] == [True] * len(w.document.pages)
    assert w.document.pages[0].setup.orientation == "landscape"
    assert w.document.pages[0].setup.size_name == "A3"
    # calculated here, across the sheets, with Excel's names
    assert [value(beams, r, 3) for r in (1, 2, 3)] == [20, 62.5, 135]
    assert value(beams, 0, 6) == 217.5
    assert value(beams, 1, 6) == "B3"
    assert value(beams, 3, 6) == pytest.approx(1.5 * 217.5)
    assert value(beams, 1, 7) == 20
    assert value(cases, 3, 2) == 3, "a dynamic array spills"
    assert value(cases, 1, 3) == pytest.approx(60 * 9.81)
    # and live: a changed load flows through both sheets
    wb = beams.sheet.workbook
    wb.set_input(beams.sheet, 1, 1, "40")
    assert value(beams, 0, 6) == pytest.approx(40 * 16 / 8 + 62.5 + 135)
    assert value(cases, 1, 3) == pytest.approx(90 * 9.81)
    assert options(beams.sheet)["titles"] == [0, 0]
    assert beams.sheet.width(0) == pytest.approx(78.75)
    assert w._last_import_report[0] == "2 worksheets came in as spreadsheet pages."
    assert any("Hidden in Excel" in line for line in w._last_import_report)


def test_insert_excel_puts_the_worksheets_after_the_page_as_one_undo_step(w, tmp_path):
    from PySide6.QtWidgets import QApplication
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    w.interactive_prompts = False
    assert w.insert_workbook(path, 0)
    QApplication.processEvents()
    assert w.document.pages[0].sheet is None
    assert all(p.sheet is not None for p in w.document.pages[1:])
    assert set(runs(w)) == {"Beams", "Load Cases"}
    # a second time: its worksheets are renamed, its formulas follow them
    assert w.insert_workbook(path, 0)
    QApplication.processEvents()
    got = runs(w)
    assert {"Beams (2)", "Load Cases (2)"} <= set(got)
    assert value(got["Beams (2)"], 3, 6) == pytest.approx(1.5 * 217.5)
    w.undo_stack.undo()
    QApplication.processEvents()
    assert set(runs(w)) == {"Beams", "Load Cases"}
    w.undo_stack.undo()
    QApplication.processEvents()
    assert runs(w) == {} and len(w.document.pages) == 1


def test_a_formula_that_differs_from_excel_is_reported(w, tmp_path):
    """Excel's last values are in the file (openpyxl writes none, so they are
    put in by hand here): a cell that comes out otherwise here is named."""
    import shutil
    import zipfile
    path = a_workbook(str(tmp_path / "beams.xlsx"))
    patched = str(tmp_path / "cached.xlsx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(patched, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                # D2 as Excel last showed it: 20 (right); D3: 99 (not what it is)
                text = text.replace("<f>B2*C2^2/8</f><v></v>", "<f>B2*C2^2/8</f><v>20</v>")
                text = text.replace("<f>B2*C2^2/8</f>", "<f>B2*C2^2/8</f><v>20</v>", 1) \
                    if "<v>20</v>" not in text else text
                text = text.replace("<f>B3*C3^2/8</f>", "<f>B3*C3^2/8</f><v>99</v>", 1)
                data = text.encode()
            out.writestr(item, data)
    shutil.copy(patched, path)
    w.interactive_prompts = False
    w.open_from_command_line(path)
    report = "\n".join(w._last_import_report)
    assert "Beams!D3" in report and "Beams!D2" not in report


def a_checks_workbook(path):
    """An engineer's check sheet: an Excel table with structured references
    of every form, conditional formatting of every kind (priorities spread
    across ranges), validation of every kind, theme colours, fit to width."""
    from openpyxl.formatting.rule import DataBarRule, FormulaRule, IconSetRule, Rule
    from openpyxl.styles.colors import Color
    from openpyxl.styles.differential import DifferentialStyle
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checks"
    rows = [("Member", "Unit Load", "Span", "DCR"), ("B1", 10, 4, 0.8), ("B2", 20, 5, 1.1),
            ("B3", 30, 6, 0.95), ("B4", 30, 2, 0.3)]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws.add_table(Table(displayName="Chk", ref="A1:D5"))
    ws["E2"] = "=Chk[@[Unit Load]]*Chk[@Span]"
    ws["F2"] = "=COUNTA(Chk[#Headers])"
    ws["F3"] = "=ROWS(Chk[#Data])"
    ws["F4"] = "=SUM(Chk[[Unit Load]:[Span]])"
    ws["F5"] = "=MAX(Chk[DCR])"
    red = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE"))
    ws.conditional_formatting.add("D2:D5", FormulaRule(formula=["D2>1"], fill=PatternFill(bgColor="FFC7CE")))
    rule = Rule(type="containsText", operator="containsText", text="B2", dxf=red)
    rule.formula = ['NOT(ISERROR(SEARCH("B2",A2)))']
    ws.conditional_formatting.add("A2:A5", rule)
    ws.conditional_formatting.add("B2:B5", Rule(type="top10", rank=1, dxf=red))
    ws.conditional_formatting.add("C2:C5", Rule(type="aboveAverage", aboveAverage=False, dxf=red))
    ws.conditional_formatting.add("D2:D5", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws.conditional_formatting.add("C2:C5", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    for kind, op, f1, f2, at in (("whole", "between", "1", "10", "H2:H5"),
                                 ("decimal", "greaterThan", "$B$2", None, "I2"),
                                 ("custom", None, "ISNUMBER(J2)", None, "J2"),
                                 ("list", None, "$A$2:$A$5", None, "K2")):
        dv = DataValidation(type=kind, operator=op, formula1=f1, formula2=f2)
        ws.add_data_validation(dv)
        dv.add(at)
    ws["A7"] = "themed"
    ws["A7"].font = Font(color=Color(theme=4, tint=0.3999))
    ws["A8"] = "shaded"
    ws["A8"].fill = PatternFill("solid", fgColor=Color(theme=5, tint=-0.25))
    ws.print_area = "A1:F5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(path)
    return path


def test_an_engineers_check_sheet_opens_as_excel_shows_it(w, tmp_path):
    from PySide6.QtWidgets import QApplication

    from calcforge.sheet import condfmt, validation
    path = a_checks_workbook(str(tmp_path / "checks.xlsx"))
    w.interactive_prompts = False
    w.open_from_command_line(path)
    QApplication.processEvents()
    run = runs(w)["Checks"]
    sheet = run.sheet
    assert [value(run, 1, 4), value(run, 1, 5), value(run, 2, 5), value(run, 3, 5), value(run, 4, 5)] == \
        [40, 4, 4, 90 + 17, 1.1]
    assert [r["type"] for r in sheet.cond_rules] == ["formula", "text", "top", "below", "databar", "icons"], \
        "in Excel's priority order across all the ranges"
    looks = condfmt.looks_for(sheet)
    fill = lambda r, c: looks.look(r, c).get("format", {}).get("fill")  # noqa: E731
    assert [fill(r, 3) for r in (1, 2, 3, 4)] == [None, "#FFC7CE", None, None], "only DCR 1.1 > 1"
    assert [fill(r, 0) for r in (1, 2, 3, 4)] == [None, "#FFC7CE", None, None], "the member B2"
    assert [fill(r, 1) for r in (1, 2, 3, 4)] == [None, None, "#FFC7CE", "#FFC7CE"], "both 30s are the top 1"
    assert [fill(r, 2) for r in (1, 2, 3, 4)] == ["#FFC7CE", None, None, "#FFC7CE"], "4 and 2 < 4.25"
    assert looks.look(4, 3).get("bar") and looks.look(1, 2).get("icon")
    assert [v["type"] for v in sheet.validations] == ["whole", "decimal", "custom", "list"]
    wb = sheet.workbook
    wb.set_input(sheet, 1, 7, "12")
    wb.set_input(sheet, 1, 8, "5")
    assert {(1, 7), (1, 8)} <= set(validation.invalid_cells(sheet)), "12 is over 10; 5 isn't over B2's 10"
    styles = wb.styles
    assert styles.get(sheet.cells[(6, 0)].style).color.upper() == "#95B3D7", "the file's accent 1, lighter 40%"
    assert styles.get(sheet.cells[(7, 0)].style).fill.upper() == "#953735", "its accent 2, darker 25%"
    assert options(sheet)["print_area"] == [0, 0, 4, 5] and options(sheet)["fit_width"]
