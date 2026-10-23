"""An imported Excel workbook is self-contained, or open to dialling in
(the user's answers by pop-up, 2026-10-19, and the user, 2026-10-24).

* Self-contained (the default): its Excel names are its own. Its cells read
  them first (a sheet's own name wins on its sheet, as in Excel), else the
  equation variable defined above the sheet; var(X) still reads the equation
  variable. It is closed as a self-contained calculation block is (the user,
  2026-10-24): equations below read only its CalcForge variables (a cell
  named in its Variable box, read bare) and any of its cells by sheet and
  cell (Seismic_Actions.H44) — never its Excel names, bare, as Sheet.Name or
  as Book.Name (the workbook's name: its file name as an identifier, unique
  in the document).
* Not self-contained: its Excel names are read from outside only dialled
  in, as a spreadsheet's own names are ("un self contained workbook variable
  names (like spreadsheet names) should not be accessible from outside
  unless I dial in like sheet1.variable ... otherwise dont show"): an
  equation's Seismic_Actions.Z or Quake_Calc.Z, a cell's ='Seismic
  Actions'!Z — never bare, so they never clash with the document's.
  Self-contained, not even dialled in ("selfcontained shows nothing and cant
  dial in").
* The choice is a tick in the import chooser, changeable later from the
  sheet tab's right-click menu and the Properties panel, one undo step,
  saved with the document. Two workbooks with the same names stay apart.
"""
from __future__ import annotations

import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialogButtonBox

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

import tests.calc.test_calcforge_window as cw  # noqa: E402
from calcforge.calc.engine.display import display_text  # noqa: E402
from calcforge.items.calc import CalcItem  # noqa: E402
from calcforge.items.sheetpage import SheetRunItem  # noqa: E402
from calcforge.sheet.store import load_sheet, sheet_to_dict  # noqa: E402
from calcforge.sheet.workbook import Workbook  # noqa: E402
from calcforge.sheet.xlsx import make_book, read_workbook, workbook_ident  # noqa: E402


def quake(path: str, z: float = 0.4) -> str:
    """Two worksheets sharing the workbook name Z; Tables has its own
    ReturnPeriodFactor; B2 reads k, which the workbook doesn't name."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seismic Actions"
    ws["A1"], ws["A2"] = z, 1.0
    ws["B2"] = "=Z*k"
    ws["B3"] = "=ReturnPeriodFactor"
    ws["B4"] = "=var(Z)"
    tables = wb.create_sheet("Tables")
    tables["A1"], tables["A2"] = 3, 1.3
    tables["B1"] = "=Z*2"
    tables["B2"] = "=ReturnPeriodFactor*2"
    wb.defined_names["Z"] = DefinedName("Z", attr_text="'Seismic Actions'!$A$1")
    wb.defined_names["ReturnPeriodFactor"] = DefinedName("ReturnPeriodFactor",
                                                         attr_text="'Seismic Actions'!$A$2")
    tables.defined_names["ReturnPeriodFactor"] = DefinedName("ReturnPeriodFactor",
                                                             attr_text="Tables!$A$2")
    wb.save(path)
    return path


# -- the engine on its own ----------------------------------------------------------------
def engine_with(*books):
    """A workbook holding each imported workbook (path, self-contained) as one."""
    wb = Workbook()
    out = []
    for path, own in books:
        got = read_workbook(path, taken={s.name for s in wb.sheets})
        make_book(got, workbook_ident(path, {b.name for b in wb.books.values()}), "", own)
        with wb.transaction("Open"):
            sheets = [wb.add_sheet(s.name) for s in got.sheets]
            for sheet, imported in zip(sheets, got.sheets):
                load_sheet(sheet, imported.data)
        out.append({s.name: s for s in sheets})
    wb.recalculate(everything=True)
    return wb, out


def test_two_self_contained_workbooks_with_the_same_names_stay_apart(tmp_path):
    a = quake(str(tmp_path / "Quake Calc.xlsx"), 0.4)
    b = quake(str(tmp_path / "Quake Calc copy.xlsx"), 0.13)
    wb, (one, two) = engine_with((a, True), (b, True))
    assert one["Tables"].cells[(0, 1)].value == pytest.approx(0.8)
    assert two["Tables (2)"].cells[(0, 1)].value == pytest.approx(0.26)
    assert sorted(b.name for b in wb.books.values()) == ["Quake_Calc", "Quake_Calc_copy"]
    # a sheet's own name wins on its sheet; the workbook's elsewhere
    assert one["Tables"].cells[(1, 1)].value == pytest.approx(2.6)
    assert one["Seismic Actions"].cells[(2, 1)].value == pytest.approx(1.0)
    # a sheet made here doesn't see a workbook's names, nor dial in to a
    # self-contained one's
    plain = wb.add_sheet("Mine")
    wb.set_input(plain, 0, 0, "=Z")
    wb.set_input(plain, 1, 0, "='Seismic Actions'!Z")
    wb.set_input(plain, 2, 0, "='Seismic Actions'!A1")
    assert plain.cells[(0, 0)].value.code == "#NAME?"
    gone = plain.cells[(1, 0)].value
    assert gone.code == "#NAME?" and "self-contained" in gone.reason
    assert plain.cells[(2, 0)].value == pytest.approx(0.4), "its cells always"
    # one not self-contained: dialled in, never bare
    wb.set_self_contained(next(b.id for b in wb.books.values() if b.name == "Quake_Calc"), False)
    assert plain.cells[(0, 0)].value.code == "#NAME?", "never bare"
    assert plain.cells[(1, 0)].value == pytest.approx(0.4)
    assert two["Tables (2)"].cells[(0, 1)].value == pytest.approx(0.26), "its own names first"


def test_a_cell_reads_its_workbooks_name_first_then_the_equation_variable(tmp_path):
    asked = []
    wb, (sheets,) = engine_with((quake(str(tmp_path / "q.xlsx")), True))

    def outside(sheet, name):
        asked.append(name)
        if name.lower() in ("k", "z"):
            return {"k": 10.0, "z": 99.0}[name.lower()]
        raise KeyError(name)
    wb.outside = outside
    wb.recalculate(everything=True)
    sa = sheets["Seismic Actions"]
    assert sa.cells[(1, 1)].value == pytest.approx(4.0), "Z the workbook's, k the equations'"
    assert sa.cells[(3, 1)].value == pytest.approx(99.0), "var(Z) is the equation variable"
    assert "k" in asked


def test_the_workbook_and_its_setting_are_saved_with_each_sheet(tmp_path):
    wb, (sheets,) = engine_with((quake(str(tmp_path / "Quake Calc.xlsx")), True))
    data = sheet_to_dict(sheets["Tables"])
    assert data["book"]["name"] == "Quake_Calc" and data["book"]["self_contained"] is True
    assert [n[:3] for n in data["names"]] == [["ReturnPeriodFactor", "Tables!$A$2", True]]
    other = Workbook()
    copy = other.add_sheet("Seismic Actions")
    load_sheet(copy, sheet_to_dict(sheets["Seismic Actions"]))
    tables = other.add_sheet("Tables")
    load_sheet(tables, data)
    other.recalculate(everything=True)
    assert tables.cells[(0, 1)].value == pytest.approx(0.8)
    assert copy.book == tables.book and other.books[copy.book].self_contained


# -- in the window --------------------------------------------------------------------------------
def pump():
    for _ in range(3):
        QApplication.instance().processEvents()


def tick(box) -> None:
    """A click on a tick box's square, where a person clicks it."""
    from PySide6.QtWidgets import QStyle, QStyleOptionButton
    option = QStyleOptionButton()
    option.initFrom(box)
    square = box.style().subElementRect(QStyle.SE_CheckBoxIndicator, option, box)
    QTest.mouseClick(box, Qt.LeftButton, Qt.NoModifier, square.center())


def results(window) -> dict:
    return {item.region.editor.root.text(): (item.region.error.message if item.region.error else
                                             display_text(item.region.display)
                                             if item.region.display else None)
            for item in window.view.scene().items()
            if isinstance(item, CalcItem) and item.region is not None}


def runs(window) -> dict:
    return {i.sheet.name: i for i in window.view.scene().items() if isinstance(i, SheetRunItem)}


def type_equations(window, page: int, lines) -> None:
    v = cw.Sheet(window)
    v.frame = window.document.pages[page].frame
    for k, text in enumerate(lines):
        v.type_at(36, 18 + 40 * k, "")
        v.keys(text)
        v.press(Qt.Key_Return)
        v.calc.leave()
        pump()


@pytest.fixture
def w(window):
    window.show()
    window.activateWindow()
    window.interactive_prompts = False
    return window


def with_workbook(w, path, k_first=True):
    """k:=10 on page 1, the workbook after it, a blank page after that."""
    if k_first:
        type_equations(w, 0, ["k:10"])
    assert w.insert_workbook(path, 0)
    pump()
    w.add_page(len(w.document.pages) - 1)
    pump()
    return len(w.document.pages) - 1


def test_equations_read_its_cells_and_variables_but_never_its_excel_names(w, tmp_path):
    last = with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")))
    type_equations(w, last, ["a:Seismic_Actions.B2", "a=", "z:Quake_Calc.Z",
                             "q:Seismic_Actions.Z", "f:Tables.ReturnPeriodFactor", "y:Z"])
    got = results(w)
    assert got["a="] == "4", "B2 = Z*k, k the equation variable above the sheets"
    for line in ("z≔Quake_Calc.Z", "q≔Seismic_Actions.Z", "f≔Tables.ReturnPeriodFactor"):
        assert "Quake_Calc is self-contained" in got[line], (line, got[line])
    assert got["y≔Z"] == "Z - not defined."
    assert any("Sheet.H44" in line and "variable" in line for line in w._last_import_report)
    # an input changed: the cells and the equations below follow; undone, back
    sheet = runs(w)["Seismic Actions"].sheet
    w.view.begin_snapshot([runs(w)["Seismic Actions"].parentItem()])
    sheet.workbook.set_input(sheet, 0, 0, "0.5")
    w.view.commit_snapshot("Edit cell")
    pump()
    assert results(w)["a="] == "5" and runs(w)["Tables"].sheet.cells[(0, 1)].value == 1.0
    w.undo_stack.undo()
    pump()
    assert results(w)["a="] == "4" and runs(w)["Tables"].sheet.cells[(0, 1)].value == 0.8


def test_the_sheet_tab_menu_changes_it_as_one_undo_step_and_it_is_saved(w, tmp_path):
    from tests.test_calc_saving import reopen, save_to
    last = with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")))
    type_equations(w, last, ["y:Seismic_Actions.Z", "y=", "b:Quake_Calc.Z", "b=",
                             "f:Tables.ReturnPeriodFactor", "f=", "n:Z"])
    assert "self-contained" in results(w)["y≔Seismic_Actions.Z"], "can't dial in"
    w.go_to_page(1)                       # the workbook view: its tabs
    pump()
    menu = w.sheet_tabs.build_menu(0)
    (tick,) = [a for a in menu.actions() if a.text() == "Self-contained"]
    assert tick.isCheckable() and tick.isChecked()
    steps = w.undo_stack.count()
    tick.trigger()
    pump()
    assert w.undo_stack.count() == steps + 1
    got = results(w)
    assert got.get("y=") == "0.4" and got.get("b=") == "0.4" and got.get("f=") == "1.3", got
    assert got["n≔Z"] == "Z - not defined.", "never bare (the user, 2026-10-24)"
    w.undo_stack.undo()
    pump()
    assert "self-contained" in results(w)["y≔Seismic_Actions.Z"]
    w.undo_stack.redo()
    pump()
    path = str(tmp_path / "doc.pdf")
    save_to(w, path)
    reopen(w, path)
    pump()
    pump()
    assert results(w)["y="] == "0.4"
    books = runs(w)["Tables"].sheet.workbook.books
    assert [(b.name, b.self_contained) for b in books.values()] == [("Quake_Calc", False)]
    w.go_to_page(1)
    pump()
    menu = w.sheet_tabs.build_menu(0)
    assert not [a for a in menu.actions() if a.text() == "Self-contained"][0].isChecked()


def test_its_thumbnails_entry_has_the_tick_too(w, tmp_path):
    """The user, 2026-10-24: "right click on workbook in thumbnail wont let
    me selection self contained vs not"."""
    with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    sheet = runs(w)["Tables"].sheet
    first = next(i for i, p in enumerate(w.document.pages) if p.sheet is not None)
    menu = w.page_menu(first)
    (act,) = [a for a in menu.actions() if a.text() == "Self-contained"]
    assert act.isCheckable() and act.isChecked()
    steps = w.undo_stack.count()
    act.trigger()
    pump()
    assert not sheet.workbook.books[sheet.book].self_contained
    assert w.undo_stack.count() == steps + 1
    assert w.document.pages[0].sheet is None
    assert not [a for a in w.page_menu(0).actions() if a.text() == "Self-contained"], \
        "a plain page has none"


def test_the_properties_panel_ticks_it_by_mouse(w, tmp_path):
    from PySide6.QtWidgets import QCheckBox
    with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    run = runs(w)["Tables"]
    w.show_properties_panel()
    w.view.tables.open(run, (0, 0))
    pump()
    box = w.properties_panel.body.findChild(QCheckBox, "workbookSelfContained")
    assert box is not None and box.isChecked() and box.isVisible()
    steps = w.undo_stack.count()
    tick(box)
    pump()
    assert not run.sheet.workbook.books[run.sheet.book].self_contained
    assert w.undo_stack.count() == steps + 1


def test_unticked_its_names_never_clash_with_the_documents(w, tmp_path, monkeypatch):
    """Z:7 above and the workbook's own Z: unticking asks nothing (its Z is
    never read bare), the equations' Z stays 7 and its own is dialled in."""
    from PySide6.QtWidgets import QMessageBox
    from calcforge.ui.workbookscope import set_self_contained
    type_equations(w, 0, ["Z:7", "k:10"])
    last = with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    sheet = runs(w)["Tables"].sheet
    w.interactive_prompts = True
    asked = []
    monkeypatch.setattr(QMessageBox, "exec", lambda box: asked.append(box.text()) or 0)
    assert set_self_contained(w, sheet.book, False)
    assert not asked and not sheet.workbook.books[sheet.book].self_contained
    w.interactive_prompts = False
    type_equations(w, last, ["y:Z", "y=", "q:Seismic_Actions.Z", "q="])
    got = results(w)
    assert got["y="] == "7" and got["q="] == "0.4"
    assert runs(w)["Seismic Actions"].sheet.cells[(1, 1)].value == pytest.approx(4.0), \
        "its cells read its own Z (and the equations' k)"


def chooser_for(w, tmp_path):
    from calcforge.ui.xlsximport import SheetChooser
    got = read_workbook(quake(str(tmp_path / "Quake Calc.xlsx")))
    chooser = SheetChooser(w, got, "Quake_Calc")
    chooser.dialog.show()
    pump()
    return chooser


def test_the_chooser_has_the_tick_and_says_what_it_means(w, tmp_path):
    from calcforge.ui.xlsximport import OPEN_SAYS, SELF_CONTAINED_SAYS
    chooser = chooser_for(w, tmp_path)
    assert chooser.self_contained.isChecked() and chooser.scope() == chooser.KEEP
    assert chooser.scope_note.text() == SELF_CONTAINED_SAYS
    tick(chooser.self_contained)
    pump()
    assert not chooser.self_contained.isChecked() and chooser.scope() == chooser.OPEN
    assert chooser.scope_note.text() == OPEN_SAYS and "never bare" in OPEN_SAYS
    QTest.mouseClick(chooser.buttons.button(QDialogButtonBox.Ok), Qt.LeftButton)
    assert chooser.dialog.result()


def untick(chooser):
    """A hand on the import chooser: Self-contained unticked, then OK."""
    chooser.dialog.show()
    tick(chooser.self_contained)
    QTest.mouseClick(chooser.buttons.button(QDialogButtonBox.Ok), Qt.LeftButton)


def test_imported_unticked_it_is_dialled_in_to(w, tmp_path):
    type_equations(w, 0, ["k:10"])
    w.drive_workbook_chooser = untick
    last = with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    sheet = runs(w)["Tables"].sheet
    assert not sheet.workbook.books[sheet.book].self_contained
    assert any("Not self-contained" in line for line in w._last_import_report)
    type_equations(w, last, ["y:Z", "q:Seismic_Actions.Z", "q="])
    got = results(w)
    assert got["y≔Z"] == "Z - not defined." and got["q="] == "0.4"


def test_a_name_dialled_in_that_isnt_let_out_says_why(w, tmp_path):
    """Sheet.Name= for a name the sheet doesn't let out stays a question
    and says why (SMath's rule made it a definition waiting for a value:
    "Fill in all empty elements")."""
    from calcforge.ui.workbookscope import set_self_contained
    path = quake(str(tmp_path / "Quake Calc.xlsx"))
    book = openpyxl.load_workbook(path)
    book["Tables"].defined_names["Kp"] = DefinedName("Kp", attr_text="Tables!$A$1")
    book.save(path)
    w.drive_workbook_chooser = untick
    last = with_workbook(w, path)
    type_equations(w, last, ["Seismic_Actions.Kp=", "Tables.Kp=", "Quake_Calc.Kp="])
    got = results(w)
    assert got["Tables.Kp="] == "3"
    assert got["Seismic_Actions.Kp="] == \
        "Kp isn't a name of Seismic Actions; Tables has one of its own: Tables.Kp."
    assert got["Quake_Calc.Kp="] == "Kp isn't a name of Quake_Calc; Tables has one of its own: Tables.Kp."
    sheet = runs(w)["Tables"].sheet
    assert set_self_contained(w, sheet.book, True)
    pump()
    assert results(w)["Tables.Kp="] == (
        "Kp: Quake_Calc is self-contained, so its Excel names are for its own cells "
        "(untick Self-contained on its sheet tab to dial in).")


def test_a_cell_named_as_a_variable_is_read_bare_below_and_after_reopening(w, tmp_path):
    """The user, 2026-10-24: "Made a cell PGA in a self contained workbook and
    this was not defined at all. Even when made global." A name made in the
    Variable box on a workbook's sheet was kept as that workbook's, so the
    equations never found it. A CalcForge variable is the document's."""
    from tests.test_calc_saving import reopen, save_to
    from calcforge.ui import calcsections
    last = with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    run = runs(w)["Seismic Actions"]
    w.view.tables.open(run, (0, 0))
    pump()
    assert w.view.tables.name_selection("PGA")
    pump()
    dn = run.sheet.workbook.names[("pga", None)]
    assert dn.variable and dn.book is None
    w.view.tables.close()
    pump()
    type_equations(w, last, ["p:PGA", "p="])
    assert results(w)["p="] == "0.4"
    # listed under its workbook: Workbooks ▸ Self-contained ▸ Quake_Calc ▸ PGA
    panel = w.variables_panel
    panel.refresh()
    book = panel.group_node("book:Quake_Calc")
    assert book is not None and book.parent() is panel.group_node("workbooks/self")
    assert [book.child(i).text(0) for i in range(book.childCount())] == ["PGA"]
    # made global (for the sections after a section break): still read
    calcsections.set_cells_global(w, run, (0, 0, 0, 0), True)
    pump()
    assert results(w)["p="] == "0.4"
    path = str(tmp_path / "doc.pdf")
    save_to(w, path)
    reopen(w, path)
    pump()
    pump()
    assert results(w)["p="] == "0.4"
    wb = runs(w)["Seismic Actions"].sheet.workbook
    assert wb.names[("pga", None)].variable


def test_a_variable_kept_with_its_workbook_by_an_older_version_comes_back_the_documents():
    wb = Workbook()
    sheet = wb.add_sheet("Seismic Actions")
    load_sheet(sheet, {"book": {"id": "book-1", "name": "Quake_Calc", "self_contained": True},
                       "cells": [[0, 0, "0.4", 0]],
                       "names": [["PGA", "'Seismic Actions'!$A$1", False, "", True],
                                 ["Z", "'Seismic Actions'!$A$1", False, "", False]]})
    assert wb.names[("pga", None)].book is None
    assert wb.names[("z", "book-1")].book == "book-1", "an Excel name stays the workbook's"


def test_variables_panel_and_name_manager_show_the_workbook(w, tmp_path):
    from calcforge.ui.names import scope_text
    from calcforge.ui.variablespanel import table_rows
    with_workbook(w, quake(str(tmp_path / "Quake Calc.xlsx")), k_first=False)
    # self-contained: its Excel names are none of the equations', so not
    # listed ("selfcontained shows nothing", the user, 2026-10-24)
    assert not [r for r in table_rows(w.document)
                if "Z" in r.name or "ReturnPeriodFactor" in r.name]
    wb = runs(w)["Tables"].sheet.workbook
    scopes = sorted(scope_text(wb, dn) for dn in wb.names.values())
    assert scopes == ["Quake_Calc (self-contained)", "Quake_Calc (self-contained)", "Tables"]
    wb.set_self_contained(runs(w)["Tables"].sheet.book, False)
    pump()
    rows = table_rows(w.document)
    # not self-contained: only as dialled in, under the workbook
    assert {r.name for r in rows if r.group} == {"Quake_Calc.Z", "Quake_Calc.ReturnPeriodFactor",
                                                 "Tables.ReturnPeriodFactor"}
    assert not [r for r in rows if not r.group and r.name in ("Z", "ReturnPeriodFactor")], \
        "never as a variable of its own"
    w.variables_panel.refresh()

    panel = w.variables_panel

    def heads():
        tree = panel.tree
        return [tree.topLevelItem(i).text(0) for i in range(tree.topLevelItemCount())]
    # Workbooks ▸ Not self-contained ▸ Quake_Calc ▸ its names as dialled in,
    # each closed to start (the user, 2026-10-23 and 2026-10-24)
    assert "Workbooks (1)" in heads() and "Z" not in heads()
    top = panel.group_node("workbooks")
    assert not top.isExpanded()
    part = panel.group_node("workbooks/open")
    assert part.text(0) == "Not self-contained (1)" and part.parent() is top
    assert panel.group_node("workbooks/self") is None, "an empty group isn't shown"
    book = panel.group_node("book:Quake_Calc")
    assert book.text(0) == "Quake_Calc (3 names)" and book.parent() is part
    assert {book.child(i).text(0) for i in range(book.childCount())} == \
        {"Quake_Calc.Z", "Quake_Calc.ReturnPeriodFactor", "Tables.ReturnPeriodFactor"}
    for head in (top, part, book):
        panel.go_to(head)                          # a click on a heading opens it
    assert panel.group_node("book:Quake_Calc").isExpanded()
    scopes = sorted(scope_text(wb, dn) for dn in wb.names.values())
    assert scopes == ["Quake_Calc", "Quake_Calc", "Tables"]


def test_a_second_import_of_the_same_workbook_is_another_workbook(w, tmp_path):
    path = quake(str(tmp_path / "Quake Calc.xlsx"))
    with_workbook(w, path, k_first=False)
    assert w.insert_workbook(path, len(w.document.pages) - 1)
    pump()
    books = runs(w)["Tables"].sheet.workbook.books
    assert sorted(b.name for b in books.values()) == ["Quake_Calc", "Quake_Calc_2"]
    # shown as its sheets are: "Quake Calc" and "Quake Calc (2)" (the user, 2026-10-23)
    from calcforge.ui.sheetpages import sections
    assert [s.name for s in sections(w.document) if s.book] == ["Quake Calc", "Quake Calc (2)"]
    assert runs(w)["Tables (2)"].sheet.cells[(0, 1)].value == pytest.approx(0.8)
