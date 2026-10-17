"""Nothing a sheet showed outside itself outlives it (the user, 2026-10-17:
an Excel workbook's formula tip stayed showing after its pages were
deleted).

A sheet or table being worked in shows things outside its own cells: the
cell's text box, the formula tip and function list, the hover tip, the
unit box on a result, and the formula bar's Name Box, formula and unit.
Whichever way the sheet goes — its pages deleted from the thumbnail menu,
the Delete Pages dialog or its sheet tab, the insert undone, its document
closed, another document's tab clicked — all of it goes too
(TableEditing.end_session, called when the sheet's item leaves the canvas
and doesn't come back, and when the canvas itself goes).
"""
from __future__ import annotations

import time

import pytest
from PySide6.QtCore import QPoint, Qt, QTimer
from PySide6.QtGui import QContextMenuEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QPushButton, QTabBar, QToolTip

from calcforge.items.sheetpage import SheetRunItem
from tests.test_sheet_pages import open_at, w  # noqa: F401
from tests.test_sheet_qa_round1 import _menus_showing, _wait_for, through_menu
from tests.test_tables import pump


def a_workbook(path) -> str:
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = "Loads"
    ws.append(["Load", "Span", "Total"])
    ws.append([10, 4, "=SUM(A2,B2)"])
    book.save(path)
    return str(path)


def the_runs(w):
    return [i for i in w.view.scene().items() if isinstance(i, SheetRunItem)]


def with_a_workbook(w, tmp_path):
    """Document ▸ Insert ▸ Excel Workbook after the first page, then C2
    picked and its formula opened with F2, the caret inside SUM( so the
    formula tip shows."""
    assert w.insert_workbook(a_workbook(tmp_path / "loads.xlsx"), 0)
    pump()
    (run,) = the_runs(w)
    open_at(w, run, "C2")
    QTest.keyClick(w.view, Qt.Key_F2)
    pump()
    te = w.view.tables
    assert te.editor is not None and te.editor.text() == "=SUM(A2,B2)"
    for _ in range(3):
        QTest.keyClick(te.editor, Qt.Key_Left)
    pump()
    assert showing(w), "the formula tip, the text box and the formula bar to start with"
    return run


def showing(w) -> list:
    """What of a sheet is still showing outside its cells."""
    te = w.view.tables
    out = []
    help_ = te._formula_help
    if help_ is not None:
        if help_.tip.isVisible():
            out.append("formula tip")
        if help_.list.isVisible():
            out.append("function list")
    if te.editor is not None and te.editor.isVisible():
        out.append("cell text box")
    box = te._result_unit
    if box is not None and box.isVisible():
        out.append("unit box")
    if QToolTip.isVisible():
        out.append("hover tip")
    bar = te.bar
    if bar is not None:
        if bar.edit.text():
            out.append("formula bar: " + bar.edit.text())
        if bar.name.text():
            out.append("Name Box: " + bar.name.text())
        if bar.unit.text():
            out.append("unit: " + bar.unit.text())
    if te._uid is not None:
        out.append("still working in it")
    return out


def gone(w) -> bool:
    return _wait_for(lambda: not showing(w), 1.0)


def sheet_page_indexes(w):
    return [i for i, p in enumerate(w.document.pages) if p.sheet is not None]


# -- the ways a sheet goes --------------------------------------------------------------------------
def test_its_pages_deleted_from_the_thumbnail_menu(w, tmp_path):
    with_a_workbook(w, tmp_path)
    (index,) = sheet_page_indexes(w)
    panel = w.pages_panel
    strip = panel.list
    centre = strip.visualItemRect(strip.item(index)).center()

    def right_click_thumbnail():
        QApplication.sendEvent(strip.viewport(), QContextMenuEvent(
            QContextMenuEvent.Mouse, centre, strip.viewport().mapToGlobal(centre)))
    found, enabled, got = through_menu(w, right_click_thumbnail, ["Delete Pages…"])
    assert found and enabled and got["clicked"], got
    assert not sheet_page_indexes(w)
    assert gone(w), showing(w)


def test_its_pages_deleted_by_the_delete_pages_dialog(w, tmp_path):
    with_a_workbook(w, tmp_path)
    (index,) = sheet_page_indexes(w)
    w.go_to_page(index)
    w.interactive_prompts = True
    clicked = []

    def press_ok():
        dialog = QApplication.activeModalWidget()
        if not isinstance(dialog, QDialog):
            QTimer.singleShot(30, press_ok)
            return
        ok = next(b for b in dialog.findChildren(QPushButton)
                  if b.text().replace("&", "") in ("OK", "Delete"))
        QTest.mouseClick(ok, Qt.LeftButton)
        clicked.append(dialog.windowTitle())
    QTimer.singleShot(30, press_ok)
    try:
        QTest.keyClick(w, Qt.Key_D, Qt.ControlModifier | Qt.ShiftModifier)   # Delete Pages…
        pump()
    finally:
        w.interactive_prompts = False
    assert clicked == ["Delete Pages"], clicked
    assert not sheet_page_indexes(w)
    assert gone(w), showing(w)


def test_its_sheet_tab_deleted(w, tmp_path):
    with_a_workbook(w, tmp_path)
    tabs = w.sheet_tabs
    tabs.refresh()
    bar = tabs.bar
    k = next(i for i in range(bar.count()) if "Loads" in bar.tabText(i))
    at = bar.tabRect(k).center()

    def right_click_tab():
        QApplication.sendEvent(bar, QContextMenuEvent(QContextMenuEvent.Mouse, at,
                                                      bar.mapToGlobal(at)))
    found, enabled, got = through_menu(w, right_click_tab, ["Delete"])
    assert found and enabled and got["clicked"], got
    assert not sheet_page_indexes(w)
    assert gone(w), showing(w)


def test_its_insert_undone(w, tmp_path):
    with_a_workbook(w, tmp_path)
    QTest.keyClick(w.view.tables.editor, Qt.Key_Escape)    # Excel: Esc leaves the cell first
    pump()
    QTest.keyClick(w.view, Qt.Key_Z, Qt.ControlModifier)
    pump()
    assert not sheet_page_indexes(w)
    assert gone(w), showing(w)
    # and the tip goes even with the formula still open in its text box
    QTest.keyClick(w.view, Qt.Key_Y, Qt.ControlModifier)
    pump()
    (run,) = the_runs(w)
    open_at(w, run, "C2")
    QTest.keyClick(w.view, Qt.Key_F2)
    for _ in range(3):
        QTest.keyClick(w.view.tables.editor, Qt.Key_Left)
    pump()
    assert "formula tip" in showing(w) or "cell text box" in showing(w)
    w.undo_stack.undo()                                  # Edit ▸ Undo, the insert
    pump()
    assert gone(w), showing(w)


def test_its_document_closed_or_another_tab_clicked(w, tmp_path):
    with_a_workbook(w, tmp_path)
    QTest.keyClick(w.view.tables.editor, Qt.Key_Escape)
    QTest.keyClick(w, Qt.Key_T, Qt.ControlModifier)    # File ▸ New Tab: a second document
    pump()
    tabs = w.document_tabs
    assert tabs.count() == 2 and tabs.currentIndex() == 1
    assert gone(w), showing(w)
    # back to it by its tab, the formula open again, then the other tab clicked
    QTest.mouseClick(tabs, Qt.LeftButton, Qt.NoModifier, tabs.tabRect(0).center())
    pump()
    (run,) = the_runs(w)
    open_at(w, run, "C2")
    QTest.keyClick(w.view, Qt.Key_F2)
    pump()
    assert showing(w)
    QTest.mouseClick(tabs, Qt.LeftButton, Qt.NoModifier, tabs.tabRect(1).center())
    pump()
    assert tabs.currentIndex() == 1
    assert gone(w), showing(w)
    # and once more, then its document closed by its tab's ×
    QTest.mouseClick(tabs, Qt.LeftButton, Qt.NoModifier, tabs.tabRect(0).center())
    pump()
    (run,) = the_runs(w)
    open_at(w, run, "C2")
    QTest.keyClick(w.view, Qt.Key_F2)
    pump()
    assert showing(w)
    close = tabs.tabButton(0, QTabBar.RightSide) or tabs.tabButton(0, QTabBar.LeftSide)
    w.confirm_discard = lambda: True
    QTest.mouseClick(close, Qt.LeftButton)
    pump()
    assert tabs.count() == 1
    assert gone(w), showing(w)
