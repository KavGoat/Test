"""An input of a big engineering workbook recalculates at once (the user,
2026-10-19: typing into the seismic workbook took up to a second; it must
feel instant), with every value what a full recalculation gives.

What makes it so (docs/SPREADSHEET_DESIGN.md, "Recalculating fast"):
formulas compiled once into closures (sheet/compiled.py), only what can have
changed worked out (a dependent whose inputs came out the same is left as it
is, volatile INDIRECTs included), and the order kept between edits.

The workbook here has the seismic workbook's shape, made with openpyxl:
column names read row by row through @, INDIRECT of a cell holding a range
name, VLOOKUP and HLOOKUP into tables, IF, powers and PI(), another sheet's
cells, and thousands of chained formulas.
"""
from __future__ import annotations

import random
import statistics
import time

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

from calcforge.sheet import formula as F  # noqa: E402
from calcforge.sheet.evaluate import evaluate_cell  # noqa: E402
from calcforge.sheet.store import load_sheet  # noqa: E402
from calcforge.sheet.values import ErrorValue, Qty  # noqa: E402
from calcforge.sheet.workbook import Workbook  # noqa: E402
from calcforge.sheet.xlsx import read_workbook  # noqa: E402

ROWS = 300            # rows of the spectrum table; 13 formulas a row


def seismic_like(path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seismic Actions"
    ws["A1"], ws["B1"] = "Z", 0.4                      # the input typed into
    ws["A2"], ws["B2"] = "Period", 1.0
    ws["A3"], ws["B3"] = "Soil", "C"
    ws["A4"], ws["B4"] = "Unrelated", 7
    ws["B5"] = "=Z*Tables!B2+R_u"
    tables = wb.create_sheet("Tables")
    tables["B1"], tables["B2"] = "factor", 1.3
    tables["N1"] = "cdtA"
    tables["P1"] = "cdtB"
    for k, soil in enumerate("ABCDE"):
        tables.cell(4 + k, 1, soil)
        tables.cell(4 + k, 2, 1.0 + k * 0.1)
        tables.cell(3, 3 + k, soil)
        tables.cell(4, 3 + k, 0.2 * (k + 1))
    for i in range(ROWS):
        r = 20 + i
        t = 0.01 * (i + 1)
        tables.cell(r, 2, t)                                        # B: period
        tables.cell(r, 10, 1.0 / (1 + t))                           # J: cdtA
        tables.cell(r, 11, 0.9 / (1 + t * t))                       # K: cdtB
        tables.cell(r, 16, "=INDIRECT($N$1)")                       # P
        tables.cell(r, 17, "=INDIRECT($P$1)")                       # Q
        tables.cell(r, 18, f"=P{r}*Z*VLOOKUP(Soil,$A$4:$B$8,2,FALSE)")
        tables.cell(r, 19, f"=9800*R{r}*B{r}^2/4/PI()^2")
        tables.cell(r, 20, f"=IF(R{r}>Q{r}*Z,R{r}-Q{r}*Z,Q{r}*Z-R{r})")
        tables.cell(r, 21, f"=IF(Period<B{r},S{r},T{r})+HLOOKUP(Soil,$C$3:$G$4,2,FALSE)")
        tables.cell(r, 22, f"=U{r}/MAX(1,ROW()-19)" if i else "=U20")
        tables.cell(r, 23, f"=V{r}+W{r - 1}" if i else "=V20")       # a running chain
        tables.cell(r, 24, f"=IFERROR(1/(W{r}-W{r}),-1)")
        tables.cell(r, 25, f"=ROUND(cdt*Z,6)")
        tables.cell(r, 26, f"=Y{r}+X{r}*0")
        tables.cell(r, 27, f"=AND(Z>0,R{r}>=0)")
        tables.cell(r, 28, f"=Unrelated_twice+0")
    tables["AD1"] = f"=SUM(W20:W{19 + ROWS})"
    tables["AD2"] = "=MAX(cdt)"
    for name, ref in (("Z", "'Seismic Actions'!$B$1"), ("Period", "'Seismic Actions'!$B$2"),
                      ("Soil", "'Seismic Actions'!$B$3"),
                      ("Unrelated_twice", "'Seismic Actions'!$B$4"),
                      ("cdtA", f"Tables!$J$20:$J${19 + ROWS}"),
                      ("cdtB", f"Tables!$K$20:$K${19 + ROWS}"),
                      ("cdt", f"Tables!$J$20:$J${19 + ROWS}"),
                      ("R_u", "Tables!$AD$1")):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wb.save(path)
    return path


def opened(path: str):
    got = read_workbook(path)
    wb = Workbook()
    with wb.transaction("Open"):
        sheets = [wb.add_sheet(s.name) for s in got.sheets]
        for sheet, imported in zip(sheets, got.sheets):
            load_sheet(sheet, imported.data)
    wb.recalculate(everything=True)
    return wb, {s.name: s for s in sheets}


def values(wb) -> dict:
    return {(s.name, k): c.value for s in wb.sheets for k, c in s.cells.items() if c.is_formula}


def same(a, b) -> bool:
    if isinstance(a, float) and isinstance(b, float):
        return a == b or (a != a and b != b)
    if isinstance(a, ErrorValue) and isinstance(b, ErrorValue):
        return a.code == b.code
    return type(a) is type(b) and a == b


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory):
    return seismic_like(str(tmp_path_factory.mktemp("speed") / "Seismic like.xlsx"))


def test_every_value_after_edits_is_a_full_recalculation(workbook_path):
    """Random edits of the inputs (and an edited formula among them): the
    values kept up edit by edit equal the workbook worked out from scratch."""
    wb, sheets = opened(workbook_path)
    sa, tables = sheets["Seismic Actions"], sheets["Tables"]
    rng = random.Random(19)
    for step in range(25):
        which = rng.choice(["Z", "Period", "Soil", "Unrelated", "factor", "cdt", "formula"])
        if which == "Z":
            wb.set_input(sa, 0, 1, str(round(rng.uniform(0.05, 0.9), 3)))
        elif which == "Period":
            wb.set_input(sa, 1, 1, str(round(rng.uniform(0.0, 4.0), 2)))
        elif which == "Soil":
            wb.set_input(sa, 2, 1, rng.choice("ABCDEF"))           # F: VLOOKUP's #N/A
        elif which == "Unrelated":
            wb.set_input(sa, 3, 1, str(rng.randint(0, 9)))
        elif which == "factor":
            wb.set_input(tables, 1, 1, str(round(rng.uniform(0.5, 2.0), 2)))
        elif which == "cdt":
            wb.set_input(tables, 0, 13, rng.choice(["cdtA", "cdtB", "nothing"]))
        else:
            row = 19 + rng.randrange(ROWS)
            wb.set_input(tables, row, 24, rng.choice(["=ROUND(cdt*Z,6)", "=Z*2", "=R20"]))
        kept = values(wb)
        wb.recalculate(everything=True)
        full = values(wb)
        bad = [k for k in full if not same(kept[k], full[k])]
        assert not bad, (step, which, bad[:5])
    # and from scratch, as saved
    from calcforge.sheet.store import sheet_to_dict
    again = Workbook()
    with again.transaction("Open"):
        copies = [again.add_sheet(s.name) for s in wb.sheets]
        for copy, sheet in zip(copies, wb.sheets):
            load_sheet(copy, sheet_to_dict(sheet))
    again.recalculate(everything=True)
    fresh = values(again)
    mine = values(wb)
    assert [k for k in fresh if not same(mine[k], fresh[k])] == []


def test_an_edit_works_out_only_what_it_changes(workbook_path, monkeypatch):
    """A cell nothing changes with: only the volatile INDIRECTs are worked
    out again (Excel works out their dependents too); Period: what reads it."""
    wb, sheets = opened(workbook_path)
    sa = sheets["Seismic Actions"]
    worked = []
    real = Workbook._compute
    monkeypatch.setattr(Workbook, "_compute", lambda self, key: worked.append(key) or real(self, key))
    formulas = len(values(wb))
    wb.set_input(sa, 9, 9, "5")                          # an empty cell, read by nothing
    assert len(worked) == 2 * ROWS, "the INDIRECTs, nothing below them"
    worked.clear()
    wb.set_input(sa, 3, 1, "8")                          # read by one column
    assert len(worked) <= 2 * ROWS + ROWS + 5
    worked.clear()
    wb.set_input(sa, 0, 1, "0.45")                       # Z: most of the workbook
    assert len(worked) > formulas // 2


FORMULAS = [
    "=A1+B1", "=A1-B1", "=A1*B1", "=A1/B1", "=A1/C1", "=A1^B1", "=C1^-1", "=(0-A1)^0.5",
    "=A1^1000", "=A1=B1", "=A1<>B1", "=A1<B1", "=A1>=B1", "=0.1+0.2=0.3", "=D1+1", "=D1&A1",
    "=E1+1", "=E1*2", "=-E1", "=-D1", "=F1*2", "=F1+A1", "=A1%", "=-A1", "=+A1", "=G1+1",
    "=IF(A1>B1,A1,B1)", "=IF(A1,,B1)", "=IF(D1,1,2)", "=IF(G1,1,2)", "=IF(A1:A3>1,1,0)",
    "=IF(C1,1)", "=IFERROR(A1/C1,\"none\")", "=IFERROR(A1:A3/C1,0)", "=IFERROR(G1,H1)",
    "=PI()*A1", "=@A1:A3", "=SUM(A1:A3)*2", "=A1:A3*2", "=MAX(A1,B1,C1)", "=Z9",
    "=H1", "=Other!A1*2", "=Other!B9", "=named*2", "=named+Other!A1", "=5 kN*A1",
    "=F1/F1", "=F1-F1", "=F1>0", "=XFD1048576+1", "=OFFSET(A1,0,0)+1", "=ROW()+A1",
    "=1/0", "=\"5\"+1", "=TRUE+1", "=A1>\"a\"", "=Q9^0.5", "=VLOOKUP(2,A1:B3,2,FALSE)",
]


@pytest.mark.parametrize("text", FORMULAS)
def test_a_compiled_formula_gives_what_walking_its_tree_gives(text):
    """Every short cut of the compiled closures against evaluate.ev, the
    reason of an error included."""
    wb = Workbook()
    sheet = wb.add_sheet("Main")
    other = wb.add_sheet("Other")
    for at, typed in (((0, 0), "3"), ((0, 1), "2"), ((0, 2), "0"), ((0, 3), "five"),
                      ((0, 4), "#N/A"), ((0, 5), "5 kN"), ((0, 6), "TRUE"), ((1, 0), "4"),
                      ((2, 0), "1"), ((1, 1), "7"), ((2, 1), "9"), ((0, 7), "=1/0"),
                      ((8, 16), "-4")):
        wb.set_input(sheet, *at, typed)
    wb.set_input(other, 0, 0, "1.5")
    wb.define_name("named", "Main!$A$2")
    row, col = 5, 9
    parsed = F.parse(text[1:], row, col)
    compiled, _r1 = evaluate_cell(wb, sheet, row, col, parsed.tree, parsed)
    walked, _r2 = evaluate_cell(wb, sheet, row, col, parsed.tree)
    if isinstance(walked, ErrorValue):
        assert isinstance(compiled, ErrorValue) and compiled.code == walked.code \
            and compiled.reason == walked.reason, (compiled, walked)
    elif isinstance(walked, Qty):
        assert isinstance(compiled, Qty) and compiled == walked and compiled.unit == walked.unit
    else:
        assert type(compiled) is type(walked) and (compiled == walked or compiled != compiled)


def test_an_input_of_a_big_workbook_recalculates_fast(workbook_path):
    """3,900 formulas behind one input: the median edit well under a
    generous bound (CPU time, so a busy machine running the suite in
    parallel slows it less). About 60 ms measured on a busy machine; the
    tree-walking evaluator that worked out every dependent took 100 ms."""
    wb, sheets = opened(workbook_path)
    sa = sheets["Seismic Actions"]
    times = []
    for k in range(9):
        t = time.process_time()
        wb.set_input(sa, 0, 1, str(0.3 + 0.01 * k))
        times.append(time.process_time() - t)
    assert statistics.median(times) < 0.25, times
    assert values(wb)[("Tables", (18 + ROWS, 22))] != 0
