"""A workbook's pictures are seen on its sheet, and a workbook's Thumbnails
menu holds what a workbook takes (the user, 2026-10-26: "the pfc excel
photos shown in the thumbnail but not in the actual sheet"; "the right click
context menu on workbook split into two columns and make markups editable
went to a right column? why?").

* A markup on a spreadsheet page is drawn as itself, over the cells: in the
  page's squares (ui/markuplayer.py) it was under the grid, which a sheet
  of white cells covers completely — and the later pages of a sheet were
  stacked under its first, whose item draws the whole grid.
* A workbook page's Thumbnails menu leaves out what a sheet doesn't take
  (Page Setup, Set Scale, Crop, Page Grid, Header/Footer, Paper Size,
  Colour Processing, a PDF's markups), its inserts and page labels in
  submenus: one column. An ordinary page keeps "Make This Page's Markups
  Editable", greyed until its PDF has markups to read. A menu taller than
  the screen scrolls; it is never laid out in columns.
"""
from __future__ import annotations

import time

import pytest
from PySide6.QtWidgets import QApplication, QStyle

openpyxl = pytest.importorskip("openpyxl")
pytest.importorskip("PIL")


def pump(seconds=0.0):
    end = time.monotonic() + seconds
    QApplication.instance().processEvents()
    while time.monotonic() < end:
        QApplication.instance().processEvents()
        time.sleep(0.01)


def white_workbook(path, tmp_path) -> str:
    """One sheet of white cells, three pages long, with a red picture at B2
    (its first page) and a blue one at B100 (a later page)."""
    from openpyxl.drawing.image import Image
    from openpyxl.styles import PatternFill
    from PIL import Image as Picture
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    white = PatternFill("solid", fgColor="FFFFFF")
    for row in range(1, 151):
        for col in range(1, 9):
            ws.cell(row, col).fill = white
    ws["A1"], ws["A150"] = "Job", "End"
    for name, colour, at in (("red", (230, 30, 30), "B2"), ("blue", (30, 30, 230), "B100")):
        picture = tmp_path / f"{name}.png"
        Picture.new("RGB", (120, 60), colour).save(picture)
        ws.add_image(Image(str(picture)), at)
    wb.save(path)
    return path


@pytest.fixture
def w(window):
    window.show()
    window.activateWindow()
    window.interactive_prompts = False
    return window


def pictures(w):
    from calcforge.items.media import ImageItem
    return sorted((i for i in w.view.scene().items() if isinstance(i, ImageItem)),
                  key=lambda i: i.sceneBoundingRect().top())


def seen_at(w, item):
    """The colour on screen in the middle of *item*, once settled."""
    w.view.centerOn(item.sceneBoundingRect().center())
    pump(0.8)                                  # past the squares' settling
    image = w.view.viewport().grab().toImage()
    pump(0.1)
    image = w.view.viewport().grab().toImage()
    middle = w.view.mapFromScene(item.mapToScene(item._rect.center()))
    ratio = image.devicePixelRatio() or 1.0
    return image.pixelColor(int(middle.x() * ratio), int(middle.y() * ratio))


def test_a_workbooks_pictures_show_over_white_cells_on_every_page(w, tmp_path):
    from calcforge.items.sheetpage import SheetRunItem
    from calcforge.ui import sheetview
    assert w.insert_workbook(white_workbook(str(tmp_path / "white.xlsx"), tmp_path), 0)
    pump(0.5)
    red, blue = pictures(w)
    run = next(i for i in w.view.scene().items() if isinstance(i, SheetRunItem))
    assert red.parentItem() is not blue.parentItem(), "on two pages of the sheet"
    for view in ("pages", "workbook view"):
        if view == "workbook view":
            sheetview.enter(w, run.uid)
            pump(0.2)
        got = seen_at(w, red)
        assert got.red() > 180 and got.green() < 90, (view, got.name())
        got = seen_at(w, blue)
        assert got.blue() > 180 and got.red() < 90, (view, got.name())
        layer = blue.parentItem().layer
        assert blue not in layer.members, "drawn as itself on a sheet's page"


def test_a_workbooks_thumbnails_menu_is_a_workbooks(w):
    w.insert_spreadsheet(0)
    pump(0.1)
    index = next(i for i, p in enumerate(w.document.pages) if p.sheet is not None)
    menu = w.page_menu(index)
    labels = [a.text() for a in menu.actions() if a.text()]
    assert labels[0] == "Sheet Layout…"
    for absent in ("Page Setup…", "Crop Pages…", "Page Grid", "Header/Footer", "Paper Size",
                   "Colour Processing…", "Make This Page's Markups Editable"):
        assert absent not in labels, absent
    assert {"Insert", "Page Labels", "Self-contained", "Print Pages…"} <= set(labels)
    inserts = next(a.menu() for a in menu.actions() if a.text() == "Insert")
    assert "Insert Excel Workbook" in [a.text() for a in inserts.actions()]
    # an ordinary page of no PDF: the entry for a PDF's markups, greyed
    ordinary = next(i for i, p in enumerate(w.document.pages) if p.sheet is None)
    entry = next(a for a in w.page_menu(ordinary).actions()
                 if a.text() == "Make This Page's Markups Editable")
    assert not entry.isEnabled()
    # a menu taller than the screen scrolls, never two columns
    assert menu.style().styleHint(QStyle.SH_Menu_Scrollable, None, menu)
