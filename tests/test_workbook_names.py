"""Excel workbooks that use named variables and ranges compute as Excel
showed them (the user, 2026-10-19: "The Excel files use named variables and
ranges so your current import falls apart").

Each construct found in the user's two engineering workbooks is reproduced
in a small workbook made with openpyxl, Excel's last values patched into the
file as Excel writes them, and every formula compared with them: numbers to
1e-9 relative, text, TRUE/FALSE and errors exactly.

* a name of a block (s_x = $J$17:$J$39) where one value is wanted takes the
  cell in the formula's own row (Excel's implicit intersection): =s_x*w,
  IF(orientation="Horizontal", …), VLOOKUP(input, …), HLOOKUP(…, row, …);
* INDIRECT of a cell holding a block's name (=INDIRECT($N$23), N23 "cdtA");
* the same names given whole where a function takes a block (SUM, VLOOKUP's
  table, INDEX, MATCH) are not cut down;
* a sheet-local name wins on its sheet over a workbook name;
* a name whose cells were deleted in Excel stays, giving #REF!;
* names spelt like Excel functions (row) or SMath units and variables (h,
  A, d, p) are the names;
* Excel's @ (_xlfn.SINGLE) and empty-text results.
"""
from __future__ import annotations

import re
import shutil
import zipfile

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

from calcforge.sheet.legacy import with_intersections  # noqa: E402
from calcforge.sheet.store import load_sheet  # noqa: E402
from calcforge.sheet.values import ErrorValue, Qty  # noqa: E402
from calcforge.sheet.workbook import Workbook  # noqa: E402
from calcforge.sheet.xlsx import read_workbook  # noqa: E402


# -- making workbooks with Excel's last values in them ------------------------------------
def with_cached(path: str, values: dict) -> str:
    """Excel's last values, put into a workbook openpyxl wrote (it writes
    none): *values* is {sheet title: {"B2": value}}; text, TRUE/FALSE and
    #errors are written with Excel's cell types, "" as Excel's <v/>."""
    with zipfile.ZipFile(path) as src:
        book = src.read("xl/workbook.xml").decode()
        titles = re.findall(r'<sheet [^>]*name="([^"]+)"', book)
        parts = {f"xl/worksheets/sheet{i + 1}.xml": t for i, t in enumerate(titles)}
        patched = path + ".tmp"
        with zipfile.ZipFile(patched, "w") as out:
            for item in src.infolist():
                data = src.read(item.filename)
                title = parts.get(item.filename)
                if title in values:
                    text = data.decode()
                    for coord, v in values[title].items():
                        text = _cache(text, coord, v)
                    data = text.encode()
                out.writestr(item, data)
    shutil.move(patched, path)
    return path


def _cache(xml: str, coord: str, v) -> str:
    m = re.search(rf'<c r="{coord}"([^>]*)><f>(.*?)</f><v>.*?</v></c>', xml)
    assert m is not None, f"no formula in {coord}"
    attrs = re.sub(r'\s*t="\w+"', "", m.group(1))
    if isinstance(v, bool):
        kind, text = "b", "1" if v else "0"
    elif isinstance(v, (int, float)):
        kind, text = "n", repr(float(v))
    elif isinstance(v, str) and v.startswith("#"):
        kind, text = "e", v
    else:
        kind, text = "str", v
    cell = f'<c r="{coord}"{attrs} t="{kind}"><f>{m.group(2)}</f>' + \
        (f"<v>{text}</v>" if text != "" else "<v/>") + "</c>"
    return xml[:m.start()] + cell + xml[m.end():]


def opened(path: str):
    """The workbook as the engine has it, without the window: each sheet
    read, put in, calculated."""
    got = read_workbook(path)
    wb = Workbook()
    with wb.transaction("Open"):
        sheets = [wb.add_sheet(s.name) for s in got.sheets]
        for sheet, imported in zip(sheets, got.sheets):
            load_sheet(sheet, imported.data)
    wb.recalculate(everything=True)
    return got, wb, {s.name: s for s in sheets}


def differences(got, sheets: dict) -> list:
    """Every formula whose value is not what Excel last showed."""
    out = []
    for imported in got.sheets:
        sheet = sheets[imported.name]
        for (r, c), excel in imported.excel_values.items():
            mine = sheet.cells[(r, c)].value
            if not same(excel, mine):
                out.append((imported.name, r, c, excel, mine))
    return out


def same(excel, mine) -> bool:
    if isinstance(mine, Qty):
        return False
    if isinstance(mine, ErrorValue):
        return excel == mine.code
    if isinstance(excel, bool) or isinstance(mine, bool):
        return type(excel) is type(mine) and excel == mine
    if isinstance(excel, (int, float)) and isinstance(mine, float):
        return excel == mine or abs(excel - mine) <= 1e-9 * abs(excel)
    return excel == mine


# -- a load take-down: names of columns used row by row ------------------------------------
def load_take_down(path: str) -> str:
    """Rows 17-19 of a load take-down: each row's inputs are read through
    names of whole columns (wt_mat = $D$17:$D$19 …), Excel cutting each to
    the formula's own row; the area A, height h and spacing p are single
    cells; row = $T$17:$T$19 says which row of a lookup table to read."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calculation Sheets"
    ws["A12"], ws["B12"] = "h", 3000
    ws["A13"], ws["B13"] = "A", 20
    ws["A14"], ws["B14"] = "p", 4000
    rows = [("Slab", "Horizontal", 2400, 200, 600, 400),
            ("Wall", "Vertical", 2400, 150, 0, 0),
            ("Deck", "Horizontal", 7850, 10, 300, 0)]
    for k, (kind, orient, wt, depth, sx, sy) in enumerate(rows):
        r = 17 + k
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"H{r}"] = kind, orient, wt, depth
        ws[f"F{r}"] = 1000
        if sx:
            ws[f"J{r}"] = sx
        if sy:
            ws[f"L{r}"] = sy
        ws[f"T{r}"] = 2 + k
        ws[f"U{r}"] = '=IF(orientation="Horizontal",wt_mat*9.81/1000,wt_mat*9.81/1000*h/1000*p/A)'
        ws[f"W{r}"] = ('=wt_mat*9.81/1000*IF(orientation="Horizontal",(IF(s_x="",0,1/(s_x/1000))'
                       '+IF(s_y="",0,1/(s_y/1000))),h/1000/(s_x/1000)/(s_y/1000))')
        ws[f"Y{r}"] = "=VLOOKUP(wt_mat,materials,2,FALSE)*w/1000*d/1000"
        ws[f"P{r}"] = '=IFERROR(HLOOKUP(input,calcs,row,FALSE),"")'
        ws[f"V{r}"] = '=IF(_xlfn.SINGLE(orientation)="Horizontal",_xlfn.SINGLE(wt_mat),0)'
        ws[f"Z{r}"] = '=IFERROR(VLOOKUP(input,input_table,2,FALSE),"")'
    ws["U16"], ws["W16"], ws["Y16"] = "Slab", "Wall", "Deck"
    ws["B43"], ws["C43"], ws["B44"], ws["C44"] = 2400, "Concrete", 7850, "Steel"
    ws["B50"], ws["C50"] = "=SUM(wt_mat)", "=INDEX(d,2)"
    ws["D50"] = "=MATCH(7850,wt_mat,0)"
    ws2 = wb.create_sheet("Data")
    ws2["A2"], ws2["B2"], ws2["A3"], ws2["B3"] = "Slab", "slab input", "Deck", "deck input"
    names = {"h": "'Calculation Sheets'!$B$12", "A": "'Calculation Sheets'!$B$13",
             "p": "'Calculation Sheets'!$B$14", "input": "'Calculation Sheets'!$B$17:$B$19",
             "orientation": "'Calculation Sheets'!$C$17:$C$19",
             "wt_mat": "'Calculation Sheets'!$D$17:$D$19", "w": "'Calculation Sheets'!$F$17:$F$19",
             "d": "'Calculation Sheets'!$H$17:$H$19", "s_x": "'Calculation Sheets'!$J$17:$J$19",
             "s_y": "'Calculation Sheets'!$L$17:$L$19", "row": "'Calculation Sheets'!$T$17:$T$19",
             "calcs": "'Calculation Sheets'!$U$16:$Y$19",
             "materials": "'Calculation Sheets'!$B$43:$C$44", "input_table": "Data!$A$2:$B$3"}
    for name, ref in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wb.save(path)
    g = 9.81 / 1000
    cached = {
        "U17": 2400 * g, "U18": 2400 * g * 3000 / 1000 * 4000 / 20, "U19": 7850 * g,
        "W17": 2400 * g * (1 / 0.6 + 1 / 0.4), "W18": "#DIV/0!", "W19": 7850 * g * (1 / 0.3),
        # VLOOKUP(wt_mat, materials, 2) is the text Concrete or Steel: times w, #VALUE!
        "Y17": "#VALUE!", "Y18": "#VALUE!", "Y19": "#VALUE!",
        "V17": 2400, "V18": 0, "V19": 7850,
        "Z17": "slab input", "Z18": "", "Z19": "deck input",
        # HLOOKUP(input, U16:Y19, row): Slab, row 2 is U17; Wall, row 3 is W18
        # (#DIV/0!, so ""); Deck, row 4 is Y19 (#VALUE!, so "")
        "P17": 2400 * g, "P18": "", "P19": "",
        "B50": 2400 + 2400 + 7850, "C50": 150, "D50": 3,
    }
    return path, cached


def test_names_of_columns_read_row_by_row_give_excels_answers(tmp_path):
    path, cached = load_take_down(str(tmp_path / "ltd.xlsx"))
    got, wb, sheets = opened(path)
    calc = sheets["Calculation Sheets"]
    # what Excel shows, worked out by hand from the inputs above
    assert calc.cells[(16, 20)].value == pytest.approx(2400 * 9.81 / 1000)
    assert calc.cells[(17, 22)].value == ErrorValue("#DIV/0!"), "Wall: s_x is empty"
    assert isinstance(calc.cells[(16, 24)].value, ErrorValue), \
        "VLOOKUP of 2400 * 'Concrete' is #VALUE!, never a spilled block"
    assert calc.cells[(16, 15)].value == pytest.approx(2400 * 9.81 / 1000)
    assert calc.cells[(17, 15)].value == ""
    assert calc.cells[(18, 25)].value == "deck input"
    assert calc.cells[(49, 1)].value == 2400 + 2400 + 7850, "SUM reads the whole block"
    assert calc.cells[(49, 2)].value == 150, "INDEX reads the whole block"
    assert calc.cells[(49, 3)].value == 3, "MATCH reads the whole block"
    with_cached(path, {"Calculation Sheets": {k: v for k, v in cached.items() if k[0] in "UVWYPZ"
                                              or k in ("B50", "C50", "D50")}})
    got, wb, sheets = opened(path)
    assert differences(got, sheets) == []


def test_the_at_sign_shows_where_excel_takes_one_cell(tmp_path):
    """Excel 365 shows a plain formula of an older workbook with @ where it
    takes one cell of a block; the formula here says the same."""
    path, _cached = load_take_down(str(tmp_path / "ltd.xlsx"))
    got = read_workbook(path)
    cells = {(e[0], e[1]): e[2] for e in got.sheets[0].data["cells"]}
    assert cells[(16, 20)] == ('=IF(@orientation="Horizontal",@wt_mat*9.81/1000,'
                               '@wt_mat*9.81/1000*h/1000*p/A)'), "h, p and A are single cells"
    assert cells[(16, 15)] == '=IFERROR(HLOOKUP(@input,calcs,@row,FALSE),"")'
    assert cells[(16, 21)] == '=IF(@(orientation)="Horizontal",@(wt_mat),0)', "Excel's own @"
    assert cells[(49, 1)] == "=SUM(wt_mat)"
    assert cells[(49, 2)] == "=INDEX(d,2)"


@pytest.mark.parametrize("text, wanted", [
    ("s_x*w", "@s_x*@w"),
    ("SUM(s_x)", "SUM(s_x)"),
    ("SUM(s_x*2)", "SUM(@s_x*2)"),                    # SUM takes a block, not an expression of one
    ("SUMPRODUCT(s_x*w)", "SUMPRODUCT(s_x*w)"),        # SUMPRODUCT works over whole arrays
    ("INDEX(s_x,2)+1", "INDEX(s_x,2)+1"),
    ("VLOOKUP(s_x,tbl,2,FALSE)", "VLOOKUP(@s_x,tbl,2,FALSE)"),
    ("COUNTIF(s_x,\">0\")", "COUNTIF(s_x,\">0\")"),
    ("SUMIFS(w,s_x,\">0\",w,s_x)", "SUMIFS(w,s_x,\">0\",w,@s_x)"),
    ("IF(s_x>0,w,0)", "IF(@s_x>0,@w,0)"),
    ("SUM(IF(TRUE,w,0))", "SUM(IF(TRUE,w,0))"),        # IF passes the block on to SUM
    ("A1:A3*2", "@A1:A3*2"),
    ("INDIRECT($N$23)", "@INDIRECT($N$23)"),
    ("SUM(OFFSET(A1,0,0,3,1))", "SUM(OFFSET(A1,0,0,3,1))"),
    ("ABS(single)", "ABS(single)"),
    ("Tables!s_x*2", "@Tables!s_x*2"),
    ("-s_x", "-@s_x"),
    ("(s_x)", "(@s_x)"),
    ("@s_x+1", "@s_x+1"),
])
def test_where_excel_takes_one_cell_of_a_block(text, wanted):
    blocks = {"s_x", "w", "tables!s_x"}
    assert with_intersections(text, lambda name: name.lower() in blocks) == wanted


def test_indirect_of_a_blocks_name_takes_the_cell_in_the_row(tmp_path):
    """The seismic workbook's spectrum table: =INDIRECT($N$23), N23 holding
    the name of a column of factors (cdtA = $J$30:$J$32), down rows 30-32;
    and the next column multiplies each row's value."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"
    ws["N23"] = "cdtA"
    for k, v in enumerate((0.54, 0.5112, 0.4554)):
        ws[f"J{30 + k}"] = v
        ws[f"B{30 + k}"] = 0.1 * (k + 1)
        ws[f"P{30 + k}"] = "=INDIRECT($N$23)"
        ws[f"Y{30 + k}"] = f"=9800*P{30 + k}*B{30 + k}^2/4/PI()^2"
    ws["J59"] = "=2*$P$30"
    wb.defined_names["cdtA"] = DefinedName("cdtA", attr_text="Tables!$J$30:$J$32")
    path = str(tmp_path / "seismic.xlsx")
    wb.save(path)
    import math
    cached = {f"P{30 + k}": v for k, v in enumerate((0.54, 0.5112, 0.4554))}
    cached.update({f"Y{30 + k}": 9800 * v * (0.1 * (k + 1)) ** 2 / 4 / math.pi ** 2
                   for k, v in enumerate((0.54, 0.5112, 0.4554))})
    cached["J59"] = 1.08
    with_cached(path, {"Tables": cached})
    got, _wb, sheets = opened(path)
    assert differences(got, sheets) == []


def test_a_sheet_local_name_wins_on_its_sheet(tmp_path):
    """ReturnPeriodFactor is local to Tables (Excel's sheet scope); a
    workbook name of the same spelling is what other sheets read."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tables"
    other = wb.create_sheet("Seismic Actions")
    ws["A1"], other["A1"] = 1.3, 0.75
    ws["B1"] = "=ReturnPeriodFactor*2"
    other["B1"] = "=ReturnPeriodFactor*2"
    other["C1"] = "=Tables!ReturnPeriodFactor*2"
    wb.defined_names["ReturnPeriodFactor"] = DefinedName(
        "ReturnPeriodFactor", attr_text="'Seismic Actions'!$A$1")
    ws.defined_names["ReturnPeriodFactor"] = DefinedName("ReturnPeriodFactor",
                                                          attr_text="Tables!$A$1")
    path = str(tmp_path / "local.xlsx")
    wb.save(path)
    with_cached(path, {"Tables": {"B1": 2.6}, "Seismic Actions": {"B1": 1.5, "C1": 2.6}})
    got, _wb, sheets = opened(path)
    assert differences(got, sheets) == []


def test_a_name_whose_cells_were_deleted_stays_and_gives_ref(tmp_path):
    """Wp = 'Seismic Actions'!#REF!: kept, as Excel keeps it, not left out."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Seismic Actions"
    ws["A1"] = "=Wp*2"
    wb.defined_names["Wp"] = DefinedName("Wp", attr_text="'Seismic Actions'!#REF!")
    path = str(tmp_path / "wp.xlsx")
    wb.save(path)
    with_cached(path, {"Seismic Actions": {"A1": "#REF!"}})
    got, wbk, sheets = opened(path)
    assert not any("Wp" in line for line in got.left_out)
    assert wbk.find_name("Wp", sheets["Seismic Actions"]) is not None
    assert differences(got, sheets) == []


def test_names_spelt_like_units_variables_and_functions_are_the_names(tmp_path):
    """h (hours), d (days), A (amperes), N (newtons) are SMath units and
    likely equation variables; ROW is an Excel function: in a workbook that
    names cells so, the names are what its formulas read."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for coord, v in (("A1", 3), ("A2", 4), ("A3", 5), ("A4", 6), ("A5", 7)):
        ws[coord] = v
    ws["B1"] = "=h*d+A*N"
    ws["B2"] = "=row*2"
    ws["B3"] = "=ROW()+row"
    for name, ref in (("h", "$A$1"), ("d", "$A$2"), ("A", "$A$3"), ("N", "$A$4"), ("row", "$A$5")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Sheet1!{ref}")
    path = str(tmp_path / "units.xlsx")
    wb.save(path)
    with_cached(path, {"Sheet1": {"B1": 3 * 4 + 5 * 6, "B2": 14, "B3": 10}})
    got, _wb, sheets = opened(path)
    assert differences(got, sheets) == []


def test_a_differing_formula_is_no_longer_reported_on_opening(window, tmp_path):
    """In the window: File ▸ Open of the load take-down says nothing differs."""
    from PySide6.QtWidgets import QApplication
    path, cached = load_take_down(str(tmp_path / "ltd.xlsx"))
    with_cached(path, {"Calculation Sheets": {k: v for k, v in cached.items()
                                              if k[0] in "UVWYPZ" or k in ("B50", "C50", "D50")}})
    window.interactive_prompts = False
    window.open_from_command_line(path)
    QApplication.processEvents()
    report = "\n".join(window._last_import_report)
    assert "another answer" not in report, report
