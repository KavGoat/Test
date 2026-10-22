"""An Excel workbook's charts and pictures come in, and what else the user
found importing a seismic parts workbook (2026-10-23: "ran into ref errors
and the charts aren't pulling through"; "all the variable names inside the
sheet were showing on the variables panel").

The workbook here is made up for the test, with the same shapes: a lookup
written INDEX(C7:G7,,MATCH(...)) (the row left out), an XY scatter of a row
of cells anchored under a table, with a title and axis titles and no legend,
a picture at the top right, and names on every input.
"""
from __future__ import annotations

import pytest
openpyxl = pytest.importorskip("openpyxl")
pytest.importorskip("PIL")

from calcforge.items.chart import ChartItem  # noqa: E402
from calcforge.items.media import ImageItem  # noqa: E402
from calcforge.items.sheetpage import SheetRunItem  # noqa: E402
from tests.test_workbook_sections import pump, w  # noqa: E402,F401

_ = w


def parts_book(path: str) -> str:
    from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.chart.layout import Layout  # noqa: F401  (openpyxl's own pieces)
    from openpyxl.drawing.image import Image
    from openpyxl.workbook.defined_name import DefinedName
    from PIL import Image as Picture

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calculations"
    ws["B2"] = 50
    for k, v in enumerate(["6months", 5, 25, 50, 100]):
        ws.cell(row=2, column=3 + k, value=v)
    for k, v in enumerate([None, 250, 250, 500, "Study Required"]):
        ws.cell(row=7, column=3 + k, value=v)
    ws["B10"] = "=INDEX(C7:G7,,MATCH(B2,C2:G2))"           # the row left out
    ws["B11"] = '=CONCATENATE("1/",B10)'
    # a row of heights and coefficients, and a scatter of them
    for k, (h, c) in enumerate([(0, 0.62), (3.4, 0.98), (6.8, 1.1), (10.2, 1.3)]):
        ws.cell(row=37, column=9 + k, value=h)
        ws.cell(row=40, column=9 + k, value=c)
    wb.defined_names["Height1"] = DefinedName("Height1", attr_text="Calculations!$I$37")
    wb.defined_names["Fph_1"] = DefinedName("Fph_1", attr_text="Calculations!$I$40")
    chart = ScatterChart()
    chart.title = "Response Coefficients vs Height"
    chart.style = 13
    chart.x_axis.title = "Coefficient (F/Wp)"
    chart.y_axis.title = "Height (m)"
    chart.legend = None
    series = Series(Reference(ws, min_col=9, max_col=12, min_row=37),
                    Reference(ws, min_col=9, max_col=12, min_row=40), title="Fph")
    series.marker.symbol = "circle"
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    chart.width, chart.height = 12, 6                       # cm
    ws.add_chart(chart, "D45")
    bars = BarChart()
    bars.add_data(Reference(ws, min_col=9, max_col=12, min_row=40), from_rows=True)
    ws.add_chart(bars, "D60")
    pie = PieChart()                                        # a kind CalcForge doesn't draw
    pie.add_data(Reference(ws, min_col=9, max_col=12, min_row=40), from_rows=True)
    ws.add_chart(pie, "M60")
    picture_path = path + ".png"
    Picture.new("RGB", (80, 40), (200, 30, 30)).save(picture_path)
    logo = Image(picture_path)
    ws.add_image(logo, "N1")
    wb.create_sheet("Tables")["A1"] = "=Calculations!B10*2"
    wb.save(path)
    return path


def imported(w, tmp_path, twice: bool = False):
    path = parts_book(str(tmp_path / "Parts Coefficient.xlsx"))
    assert w.insert_workbook(path, 0)
    pump()
    if twice:
        assert w.insert_workbook(path, len(w.document.pages) - 1)
        pump()
    pump()
    return path


def runs(w) -> dict:
    return {i.sheet.name: i for i in w.view.scene().items() if isinstance(i, SheetRunItem)}


def markups(w, kind) -> list:
    return [m for f in w.scene.frames for m in f.markups() if isinstance(m, kind)]


def test_index_with_the_row_left_out_reads_as_excel(w, tmp_path):
    imported(w, tmp_path)
    calc = runs(w)["Calculations"].sheet
    assert calc.cells[(9, 1)].value == 500             # INDEX(C7:G7,,4)
    assert calc.cells[(10, 1)].value == "1/500"
    assert runs(w)["Tables"].sheet.cells[(0, 0)].value == 1000
    assert not [line for line in w._last_import_report if "another answer" in line]


def test_the_charts_come_in_where_excel_puts_them(w, tmp_path):
    imported(w, tmp_path)
    charts = markups(w, ChartItem)
    kinds = sorted(c.spec["kind"] for c in charts)
    assert kinds == ["column", "scatter"], "the pie is left out, and said so"
    assert any("1 chart of a kind" in line for line in w._last_import_report)
    scatter = next(c for c in charts if c.spec["kind"] == "scatter")
    spec = scatter.spec
    assert spec["title"] == "Response Coefficients vs Height"
    assert (spec["x_title"], spec["y_title"]) == ("Coefficient (F/Wp)", "Height (m)")
    assert spec["legend"] == "none" and spec["markers"] and not spec["lines"]
    (series,) = spec["series"]
    # the coefficients along x, the heights up y, as in the user's chart
    assert series["x"].endswith("!$I$40:$L$40") and series["y"].endswith("!$I$37:$L$37")
    assert series["name"] == "Fph"
    data = scatter.data()[0]
    assert data.xs == [0.62, 0.98, 1.1, 1.3] and data.ys == [0, 3.4, 6.8, 10.2]
    # at D45, as big as Excel shows it (12 × 6 cm)
    run = runs(w)["Calculations"]

    def at(item, row, col):
        corner = run.mapFromScene(item.scenePos())
        cell = run.cell_rect(row, col, False).topLeft()
        return abs(corner.x() - cell.x()) < 0.5 and abs(corner.y() - cell.y()) < 0.5
    assert at(scatter, 44, 3)
    assert scatter.rect.width() == pytest.approx(12 * 72 / 2.54, rel=0.02)
    assert scatter.rect.height() == pytest.approx(6 * 72 / 2.54, rel=0.02)
    # it moves with its cells
    from calcforge.ui import sheetview
    sheetview.enter(w, run.uid)
    pump()
    w.view.tables.open(run)
    w.view.tables.select((2, 0), (4, 0))        # three rows above it
    w.view.tables.insert_rows()
    pump()
    run = runs(w)["Calculations"]
    assert at(scatter, 47, 3)


def test_the_picture_comes_in(w, tmp_path):
    imported(w, tmp_path)
    (picture,) = markups(w, ImageItem)
    assert picture.pixmap() is not None and not picture.pixmap().isNull()
    assert w.document.asset(picture.asset_key)
    run = runs(w)["Calculations"]
    corner = run.mapFromScene(picture.scenePos())
    cell = run.cell_rect(0, 13, False).topLeft()                      # N1
    assert abs(corner.x() - cell.x()) < 0.5 and abs(corner.y() - cell.y()) < 0.5
    assert picture.local_rect().width() == pytest.approx(80 * 0.75, rel=0.05)
    assert not [line for line in w._last_import_report if "Left out" in line and "picture" in line]


def test_imported_twice_each_copy_reads_its_own_and_is_named_2(w, tmp_path):
    imported(w, tmp_path, twice=True)
    from calcforge.ui.sheetpages import sections
    assert [s.name for s in sections(w.document) if s.book] == [
        "Parts Coefficient", "Parts Coefficient (2)"]
    charts = [c for c in markups(w, ChartItem) if c.spec["kind"] == "scatter"]
    assert len(charts) == 2
    second = [c for c in charts if "(2)" in c.spec["series"][0]["x"]]
    assert len(second) == 1, "the copy's chart reads the copy's cells"
    assert runs(w)["Calculations (2)"].sheet.cells[(9, 1)].value == 500


def test_the_variables_panel_lists_calcforge_variables_only(w, tmp_path):
    imported(w, tmp_path)
    panel = w.variables_panel
    panel.refresh()

    def names():
        return [panel.tree.topLevelItem(i).text(0) for i in range(panel.tree.topLevelItemCount())]
    assert not [n for n in names() if "Height1" in n or "Fph_1" in n]
    panel.sheet_names.setChecked(True)
    assert [n for n in names() if n.endswith("Height1")] and [n for n in names() if n.endswith("Fph_1")]
    panel.sheet_names.setChecked(False)
    # a cell named with its Variable box is a CalcForge variable: listed
    calc = runs(w)["Calculations"].sheet
    calc.workbook.define_name("W_total", "Calculations!$I$40", variable=True)
    panel.refresh()
    assert [n for n in names() if n.endswith("W_total")], "listed with Sheet Names unticked"
