"""A designed workbook comes in as its designer saw it in Excel (the user,
2026-10-29): its conditional formats from Excel's 2010 extension list too
(a rule reading another sheet goes there, and openpyxl drops it), its
superscripts and subscripts drawn raised and lowered, and a print area that
runs over the page by a hair at Excel's "Adjust to N %" on one page, not
with its last row on a page of its own."""
from __future__ import annotations

import zipfile

import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402

from calcforge.items.sheetpage import SheetRunItem  # noqa: E402
from calcforge.sheet.xlsx import read_workbook  # noqa: E402

X14 = ('<extLst><ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" '
       'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
       '<x14:conditionalFormattings>'
       '<x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
       '<x14:cfRule type="expression" priority="1"><xm:f>OR($A$1=Data!$A$1,$A$1=Data!$A$2)</xm:f>'
       '<x14:dxf><fill><patternFill><bgColor rgb="FFDCE6F2"/></patternFill></fill></x14:dxf>'
       '</x14:cfRule><xm:sqref>B1:B3</xm:sqref></x14:conditionalFormatting>'
       '</x14:conditionalFormattings></ext></extLst>')


def pump(seconds=0.0):
    import time
    end = time.monotonic() + seconds
    QApplication.processEvents()
    while time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)


def with_extension_rule(path) -> str:
    """A workbook whose Inputs sheet has a conditional format Excel keeps
    in its extension list: B1:B3 blue while A1 names a material of Data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "steel"
    ws["B1"], ws["B2"], ws["B3"] = 1, 2, 3
    data = wb.create_sheet("Data")
    data["A1"], data["A2"] = "timber", "concrete"
    wb.save(str(path))
    # Excel writes the rule in the sheet's extLst; openpyxl never does
    with zipfile.ZipFile(str(path)) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf8")
    assert "</worksheet>" in sheet
    parts["xl/worksheets/sheet1.xml"] = sheet.replace("</worksheet>", X14 + "</worksheet>").encode("utf8")
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in parts.items():
            z.writestr(name, body)
    return str(path)


def test_a_conditional_format_from_excels_extension_list_comes_in(tmp_path):
    got = read_workbook(with_extension_rule(tmp_path / "ext.xlsx"))
    inputs = next(s for s in got.sheets if s.name == "Inputs")
    rules = inputs.data["cond_rules"]
    assert len(rules) == 1
    assert rules[0]["type"] == "formula"
    assert rules[0]["formula"] == "=OR($A$1=Data!$A$1,$A$1=Data!$A$2)"
    assert rules[0]["ranges"] == [[0, 1, 2, 1]]
    assert rules[0]["format"] == {"fill": "#DCE6F2"}


def runs_of(w) -> dict:
    return {i.sheet.name: i for i in w.view.scene().items() if isinstance(i, SheetRunItem)}


def test_the_extension_rule_colours_the_cells_as_the_choice_changes(window, tmp_path):
    w = window
    w.show()
    w.interactive_prompts = False
    w.insert_workbook(with_extension_rule(tmp_path / "ext.xlsx"), 0)
    pump(0.3)
    run = runs_of(w)["Inputs"]
    sheet, wb = run.sheet, run.sheet.workbook
    assert run._shown_style(0, 1).fill is None, "steel is no material of Data: plain"
    wb.set_input(sheet, 0, 0, "timber")
    pump(0.2)
    assert run._shown_style(0, 1).fill == "#DCE6F2"
    assert run._shown_style(2, 1).fill == "#DCE6F2"
    assert run._shown_style(0, 2).fill is None, "only the rule's range"


def test_excels_superscripts_and_subscripts_are_drawn_raised_and_lowered(window, tmp_path):
    path = str(tmp_path / "rich.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Units"
    ws["A1"] = CellRichText("kN/m", TextBlock(InlineFont(vertAlign="superscript"), "3"))
    ws["A2"] = CellRichText("f'", TextBlock(InlineFont(vertAlign="subscript"), "c"), " = 40 MPa")
    ws["A3"] = CellRichText("plain ", TextBlock(InlineFont(b=True), "bold"))
    wb.save(path)
    got = read_workbook(path)
    units = got.sheets[0]
    by_cell = {(e[0], e[1]): (e[2], units.data["styles"][e[3]]) for e in units.data["cells"]}
    assert by_cell[(0, 0)][0] == "kN/m3"
    assert by_cell[(0, 0)][1]["runs"] == [["kN/m", ""], ["3", "up"]]
    assert by_cell[(1, 0)][1]["runs"] == [["f'", ""], ["c", "down"], [" = 40 MPa", ""]]
    assert "runs" not in by_cell[(2, 0)][1], "bold words alone are plain text"
    # on the page: the drawn line has its raised run, in the cell's ink
    from calcforge.items.table import Drawn, rich_runs, _line, _metrics
    w = window
    w.show()
    w.interactive_prompts = False
    w.insert_workbook(path, 0)
    pump(0.3)
    run = runs_of(w)["Units"]
    st = run._shown_style(0, 0)
    assert st.runs == (("kN/m", ""), ("3", "up"))
    line = _line("kN/m3", None, st, _metrics(st))
    assert line.rich and [k for _t, k in line.runs] == ["", "up"]
    plain = Drawn("kN/m3", None, line.font, line.metrics)
    assert line.width < plain.width, "the 3 is smaller"
    assert line.ascent > plain.ascent, "and raised"
    below = run._shown_style(1, 0)
    low = _line("f'c = 40 MPa", None, below, _metrics(below))
    assert low.descent > Drawn("f'c = 40 MPa", None, low.font, low.metrics).descent, "lowered"
    # retyped, the text is plain again: the runs no longer spell it
    assert rich_runs(st, "kN/m2") is None
    run.sheet.workbook.set_input(run.sheet, 0, 0, "kN/m2")
    pump(0.1)
    assert not _line("kN/m2", None, run._shown_style(0, 0), _metrics(st)).rich


def designed(path, scale: int, rows: int) -> str:
    """A4 landscape at Excel's Adjust to *scale* %, a print area of *rows*
    rows of 14 pt and 18 columns, the Load Take Down sheet's shape."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    for r in range(1, rows + 1):
        ws.cell(r, 1, f"row {r}")
        ws.row_dimensions[r].height = 14.0
    ws.print_area = f"A1:R{rows}"
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.scale = scale
    ws.page_margins.left = ws.page_margins.right = 0.236
    ws.page_margins.top = ws.page_margins.bottom = 0.748
    wb.save(str(path))
    return str(path)


def test_a_print_area_a_hair_over_the_page_at_its_scale_fits_the_page(window, tmp_path):
    # 47 rows at 76 %: 500 pt of rows on 487.6 pt of page — Excel's own
    # page breaks fit them; here the scale comes down to the 74 % that does
    got = read_workbook(designed(tmp_path / "ltd.xlsx", 76, 47))
    page = got.sheets[0].data["page"]
    assert page["scale"] == 74 and not page.get("fit_width") and not page.get("fit_tall")
    w = window
    w.show()
    w.interactive_prompts = False
    w.insert_workbook(str(tmp_path / "ltd.xlsx"), 0)
    pump(0.3)
    landscape = [p for p in w.document.pages if p.setup.orientation == "landscape"]
    assert len(landscape) == 1, "one page, as designed"


def test_a_print_area_well_over_the_page_keeps_its_scale(tmp_path):
    got = read_workbook(designed(tmp_path / "long.xlsx", 76, 60))
    assert got.sheets[0].data["page"]["scale"] == 76, "60 rows are two pages: left alone"
    got = read_workbook(designed(tmp_path / "fits.xlsx", 76, 40))
    assert got.sheets[0].data["page"]["scale"] == 76, "40 rows fit: left alone"


def test_delete_over_the_thumbnails_deletes_the_pages_even_with_a_markup_selected(window):
    """The window's Delete shortcut (a selected markup) took the key before
    the Thumbnails strip saw it, and the pages picked out stayed (the user,
    2026-10-29: "still can't click delete keyboard to delete multiple
    selected pages")."""
    from PySide6.QtCore import QRectF
    from calcforge.items.shapes import RectItem
    w = window
    w.show()
    w.interactive_prompts = False
    for _ in range(3):
        w.quick_add_page(len(w.document.pages) - 1)
    pump(0.2)
    frame = w.scene.frames[0]
    box = frame.add_markup(RectItem("rect", QRectF(0, 0, 80, 50)))
    box.setSelected(True)
    pump(0.1)
    assert w.act_delete.isEnabled(), "Delete Markup is live"
    strip = w.pages_panel.list
    strip.clearSelection()
    for row in (1, 2):
        strip.item(row).setSelected(True)
    strip.setCurrentRow(2)
    strip.setFocus(Qt.OtherFocusReason)
    pump(0.1)
    QTest.keyClick(strip, Qt.Key_Delete)
    pump(0.3)
    assert len(w.document.pages) == 2, "the two pages picked out went"
    assert box.scene() is not None and box in frame.built_markups(), "the markup stayed"


def test_an_ordinary_pages_menu_is_shaped_like_a_workbook_pages(window, tmp_path):
    """Every kind of page under one Insert, the labels under Page Labels,
    on a PDF's page as on a workbook's (the user, 2026-10-29: "make things
    consistent throughout the whole app")."""
    w = window
    w.show()
    w.interactive_prompts = False
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(str(tmp_path / "one.xlsx"))
    w.insert_workbook(str(tmp_path / "one.xlsx"), 0)
    pump(0.2)
    plain = next(i for i, p in enumerate(w.document.pages) if p.sheet is None)
    book = next(i for i, p in enumerate(w.document.pages) if p.sheet is not None)

    def shape(index):
        menu = w.page_menu(index)
        return [(a.text(), a.menu() is not None) for a in menu.actions()
                if not a.isSeparator() and a.isVisible()]
    ordinary, workbook = shape(plain), shape(book)
    assert ("Insert", True) in ordinary and ("Page Labels", True) in ordinary
    assert ("Insert", True) in workbook and ("Page Labels", True) in workbook
    assert not any(text.startswith("Insert ") for text, _sub in ordinary), \
        "no kind of page loose at the top level"
    assert not any(text.endswith("Page Label") for text, _sub in ordinary)
    # the same entries in the same order, but for what only a PDF page has
    only_a_pdf = {"Page Setup…", "Set Scale…", "Crop Pages", "Colour Processing…",
                  "Color Processing…", "Page Grid", "Header/Footer", "Paper Size",
                  "Make This Page's Markups Editable"}
    shared = [entry for entry in ordinary if entry[0] not in only_a_pdf]
    book_shared = [entry for entry in workbook
                   if entry[0] not in ("Sheet Layout…",) and not entry[0].startswith("Self")]
    assert [e for e in shared if e in book_shared] == [e for e in book_shared if e in shared]
